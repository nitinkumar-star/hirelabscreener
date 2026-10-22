"""
Profit & Cash — modules/finance.py

Run:   pytest test_finance.py -q
"""
import io
import os
import sys
import tempfile
import datetime

import pytest

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)

CO, OTHER = 997, 996
ADMIN, REC, FL, OADMIN = 9971, 9972, 9973, 9961
TODAY = (datetime.datetime.utcnow() + datetime.timedelta(hours=5, minutes=30)).date()


def d(days=0):
    return (TODAY + datetime.timedelta(days=days)).isoformat()


def last_month_day():
    return (TODAY.replace(day=1) - datetime.timedelta(days=3)).isoformat()


@pytest.fixture(scope='module')
def srv():
    os.environ.setdefault('DATA_DIR', tempfile.mkdtemp(prefix='fin_'))
    import contextlib
    with contextlib.redirect_stdout(io.StringIO()):
        import server
    server.app.config['TESTING'] = True
    conn = server.get_db(); c = conn.cursor(); now = server.ts()

    def ins(t, **kw):
        c.execute(f"INSERT OR REPLACE INTO {t} ({','.join(kw)}) VALUES ({','.join('?' * len(kw))})", list(kw.values()))
    ins('companies', id=CO, name='Fin Co', status='active', created_at=now)
    ins('companies', id=OTHER, name='Other', status='active', created_at=now)
    for uid, un, co, adm, role in ((ADMIN, 'fin_admin', CO, 1, 'user'), (REC, 'fin_rec', CO, 0, 'user'),
                                   (FL, 'fin_fl', CO, 0, 'freelancer_sourcer'), (OADMIN, 'fin_o', OTHER, 1, 'user')):
        ins('users', id=uid, username=un, password_hash='x', display_name=un.title(), role=role, status='approved',
            company_id=co, is_company_admin=adm, created_at=now)
    ins('crm_clients', id=99701, company_id=CO, name='Beta Pvt Ltd', created_at=now)
    ins('crm_contacts', id=99702, company_id=CO, client_id=99701, name='Accounts', email='ap@beta.test', is_primary=1, created_at=now)
    ins('mandates', id=99711, role='Substation PM', client='Acme Ltd', status='active', owner_id=CO, created_at=now)
    ins('mandates', id=99712, role='Site Engineer', client='Beta Pvt Ltd', status='active', owner_id=CO, created_at=now)
    # candidate placed by the freelancer and invoiced; another joined but never invoiced
    ins('candidates', id=99721, mandate_id=99711, name='Ravi Kumar', stage='Joined', owner_id=CO, created_at=now, sourced_by=FL)
    ins('candidates', id=99722, mandate_id=99712, name='Sita Iyer', stage='Joined', owner_id=CO, created_at=now,
        offered_ctc=1200000, fee_percent=8.33, joining_date=d(-5))
    ins('candidates', id=99723, mandate_id=99712, name='Dropped Guy', stage='Offer Dropped', owner_id=CO, created_at=now,
        offered_ctc=900000, joining_date=d(-5))
    # invoices: A paid this month (full, TDS cut), B unpaid & overdue, C proforma (ignored), D short-paid last month
    ins('invoices', id=99731, owner_id=CO, invoice_no='HL/001', invoice_date=d(0), buyer_name='Acme Ltd', amount=300000, gst_rate=18,
        status='Paid', received_amount=0, received_date=d(0), candidate_id=99721, mandate_id=99711, candidate_name='Ravi Kumar',
        role='Substation PM', doc_type='tax', created_at=now)
    ins('invoices', id=99732, owner_id=CO, invoice_no='HL/002', invoice_date=d(-70), due_date=d(-40), buyer_name='Beta Pvt Ltd',
        amount=200000, gst_rate=18, status='Sent', client_id=99701, doc_type='tax', created_at=now)
    ins('invoices', id=99733, owner_id=CO, invoice_no='PF/001', invoice_date=d(0), buyer_name='Acme Ltd', amount=999999,
        gst_rate=18, status='Sent', doc_type='proforma', created_at=now)
    ins('invoices', id=99734, owner_id=CO, invoice_no='HL/000', invoice_date=last_month_day(), buyer_name='ACME LIMITED',
        amount=100000, gst_rate=18, status='Paid', received_amount=100000, received_date=last_month_day(), doc_type='tax', created_at=now)
    ins('invoices', id=99741, owner_id=OTHER, invoice_no='X/1', invoice_date=d(0), buyer_name='Secret', amount=5000000,
        gst_rate=18, status='Paid', doc_type='tax', created_at=now)
    # expenses this month: Naukri with GST (direct), rent (operating)
    ins('expenses', id=99751, owner_id=CO, date=d(0), category='Job Boards & Sourcing', payee='Naukri', amount=59000,
        gst_amount=9000, note='', invoice_id=0, created_at=now)
    ins('expenses', id=99752, owner_id=CO, date=d(0), category='Rent', payee='Landlord', amount=50000, note='', invoice_id=0, created_at=now)
    conn.commit(); conn.close()
    yield server


def cl(server, uid):
    c = server.app.test_client()
    with c.session_transaction() as s:
        s['user_id'] = uid
    return c


def test_admin_only_and_tenant(srv):
    assert cl(srv, REC).get('/api/finance/summary').status_code == 403
    assert cl(srv, FL).get('/api/finance/summary').status_code in (401, 403)
    o = cl(srv, OADMIN).get('/api/finance/summary').get_json()
    assert o['ok'] and o['pnl']['revenue'] == 5000000          # only its own invoice


def test_pnl_without_gst_with_commission(srv):
    a = cl(srv, ADMIN)
    # Beta deducts 2% TDS (set before the numbers are read)
    assert a.put('/api/finance/setup', json={'clients': [{'client': 'Beta Pvt Ltd', 'tds_percent': 2}]}).get_json()['ok']
    j = a.get('/api/finance/summary?period=this_month').get_json()
    p = j['pnl']
    assert j['settings']['tds_percent'] == 10 and j['settings']['freelancer_percent'] == 15
    assert p['revenue'] == 300000                                # taxable value only, no GST, proforma ignored
    assert p['gst_collected'] == 54000 and p['tds'] == 30000
    # direct: Naukri 59,000 − 9,000 GST = 50,000 + auto commission 15% × 3,00,000 = 45,000
    assert p['direct_costs'] == 95000 and p['operating_costs'] == 50000
    assert p['gross_profit'] == 205000 and p['net_profit'] == 155000 and p['net_margin'] == 51.7
    assert p['gst_on_expenses'] == 9000 and p['unpaid_expenses'] == 45000
    # cash view: received 3,54,000 − 30,000 TDS; paid out Naukri + rent (commission still unpaid)
    c = j['cash']
    assert c['money_in'] == 324000 and c['money_out'] == 109000 and c['net_cash'] == 215000
    # target
    t = j['target']
    assert t['target'] == 3000000 and t['billed'] == 300000 and t['pct'] == 10.0 and t['gap'] == 2700000
    assert t['placements_needed'] >= 1
    # commission row exists once, linked to the freelancer and candidate
    conn = srv.get_db()
    rows = conn.execute("SELECT * FROM expenses WHERE owner_id=? AND source_key LIKE 'flc:%'", (CO,)).fetchall()
    conn.close()
    assert len(rows) == 1 and rows[0]['amount'] == 45000 and rows[0]['freelancer_user_id'] == FL and rows[0]['status'] == 'unpaid'
    a.get('/api/finance/summary')
    conn = srv.get_db()
    assert conn.execute("SELECT COUNT(*) FROM expenses WHERE owner_id=? AND source_key LIKE 'flc:%'", (CO,)).fetchone()[0] == 1
    conn.close()


def test_collections_tds_gst_leakage(srv):
    j = cl(srv, ADMIN).get('/api/finance/summary?period=fy').get_json()
    ag = j['aging']
    # Beta: 2,36,000 − 2% TDS (4,000) = 2,32,000, 70 days old, 40 days late
    assert ag['total'] == 232000 and ag['buckets']['61-90'] == 232000
    b = ag['invoices'][0]
    assert b['no'] == 'HL/002' and b['overdue_days'] == 40
    # Acme and ACME LIMITED are the same client
    tds = {r['client_key']: r for r in j['tds']}
    assert tds['acme']['invoiced'] >= 300000 and tds['beta']['tds_percent'] == 2
    # short paid: 1,18,000 − 10,000 TDS = 1,08,000 expected, 1,00,000 received
    assert any(s['no'] == 'HL/000' and s['short'] == 8000 for s in j['short_paid'])
    # joined, not invoiced: Sita (8.33% of 12 L), not the dropped candidate, not Ravi (invoiced)
    ni = j['not_invoiced']
    assert [x['name'] for x in ni] == ['Sita Iyer'] and ni[0]['expected_fee'] == 99960
    # GST month table for this month: 54,000 out, 9,000 in
    gm = [g for g in j['gst'] if g['month'] == TODAY.strftime('%b %Y')][0]
    assert gm['output_gst'] == 54000 and gm['input_gst'] == 9000 and gm['payable'] == 45000
    # forecast includes the overdue Beta money now
    assert j['forecast']['total_in'] >= 232000
    cl_rows = {r['client_key']: r for r in j['clients']}
    assert cl_rows['beta']['overdue'] == 232000 and cl_rows['acme']['avg_pay_days'] == 0
    assert j['cost_per_hire']['job_board_spend'] == 50000


def test_expense_crud_recurring_budget(srv):
    a = cl(srv, ADMIN)
    assert a.post('/api/finance/expenses', json={'amount': 100, 'gst_amount': 200}).status_code == 400
    start = (TODAY.replace(day=1) - datetime.timedelta(days=40)).replace(day=1)   # two months back
    r = a.post('/api/finance/expenses', json={'date': start.isoformat(), 'category': 'Tools & Subscriptions', 'payee': 'Zoho',
                                              'amount': 1180, 'gst_amount': 180, 'recurring': True}).get_json()
    assert r['ok']
    a.get('/api/finance/summary'); a.get('/api/finance/summary')
    conn = srv.get_db()
    n = conn.execute("SELECT COUNT(*) FROM expenses WHERE owner_id=? AND payee='Zoho'", (CO,)).fetchone()[0]
    conn.close()
    assert n == 3                                   # start month + 2 more up to this month, never duplicated
    eid = r['id']
    assert a.put('/api/finance/expense/%d' % eid, json={'note': 'CRM', 'status': 'unpaid'}).get_json()['ok']
    lst = a.get('/api/finance/expenses?all=1').get_json()['expenses']
    e = [x for x in lst if x['id'] == eid][0]
    assert e['note'] == 'CRM' and e['status'] == 'unpaid' and e['amount'] == 1180
    assert a.put('/api/finance/expense/%d' % eid, json={'status': 'paid'}).get_json()['ok']
    # budget alert
    a.put('/api/finance/setup', json={'budgets': [{'category': 'Rent', 'monthly_amount': 55000}]})
    b = {x['category']: x for x in a.get('/api/finance/summary').get_json()['budgets']}
    assert b['Rent']['state'] == 'near' and b['Rent']['pct'] == 91
    # other company cannot touch it
    assert cl(srv, OADMIN).delete('/api/finance/expense/%d' % eid).status_code == 404


def test_deleted_commission_stays_deleted_and_rates(srv):
    a = cl(srv, ADMIN)
    conn = srv.get_db()
    cid = conn.execute("SELECT id FROM expenses WHERE owner_id=? AND source_key='flc:99721'", (CO,)).fetchone()[0]
    conn.close()
    assert a.delete('/api/finance/expense/%d' % cid).get_json()['ok']
    a.get('/api/finance/summary')
    conn = srv.get_db()
    rows = conn.execute("SELECT amount FROM expenses WHERE owner_id=? AND source_key='flc:99721'", (CO,)).fetchall()
    conn.close()
    assert len(rows) == 1 and rows[0][0] == 0
    # per-freelancer fixed rate is stored; blank goes back to the default
    a.put('/api/finance/setup', json={'freelancers': [{'id': FL, 'mode': 'fixed', 'value': 25000}]})
    f = [x for x in a.get('/api/finance/setup').get_json()['freelancers'] if x['id'] == FL][0]
    assert f['mode'] == 'fixed' and f['value'] == 25000
    a.put('/api/finance/setup', json={'freelancers': [{'id': FL, 'value': ''}]})
    f = [x for x in a.get('/api/finance/setup').get_json()['freelancers'] if x['id'] == FL][0]
    assert f['value'] is None
    # a user of another company cannot be given a rate
    a.put('/api/finance/setup', json={'freelancers': [{'id': OADMIN, 'value': 5}]})
    conn = srv.get_db()
    assert not conn.execute('SELECT 1 FROM finance_freelancer_rate WHERE user_id=?', (OADMIN,)).fetchone()
    conn.close()


def test_old_placements_get_no_commission(srv):
    """A freelancer placement invoiced before the commission start date is not added."""
    conn = srv.get_db()
    conn.execute("INSERT OR REPLACE INTO candidates (id, mandate_id, name, stage, owner_id, created_at, sourced_by) VALUES (99724, 99711, 'Old One', 'Joined', ?, ?, ?)", (CO, srv.ts(), FL))
    conn.execute("INSERT OR REPLACE INTO invoices (id, owner_id, invoice_no, invoice_date, buyer_name, amount, gst_rate, status, candidate_id, doc_type, created_at) "
                 "VALUES (99735, ?, 'HL/OLD', ?, 'Acme Ltd', 150000, 18, 'Paid', 99724, 'tax', ?)", (CO, d(-200), srv.ts()))
    conn.commit(); conn.close()
    cl(srv, ADMIN).get('/api/finance/summary')
    conn = srv.get_db()
    assert not conn.execute("SELECT 1 FROM expenses WHERE source_key='flc:99724'").fetchone()
    conn.close()
    # moving the start date back adds it (15% of 1,50,000)
    a = cl(srv, ADMIN)
    a.put('/api/finance/setup', json={'settings': {'commission_from': d(-365)}})
    a.get('/api/finance/summary')
    conn = srv.get_db()
    r = conn.execute("SELECT amount FROM expenses WHERE source_key='flc:99724'").fetchone()
    conn.close()
    assert r and r[0] == 22500


def test_reminder(srv, monkeypatch):
    a = cl(srv, ADMIN)
    dr = a.get('/api/finance/reminder?invoice_id=99732').get_json()
    assert dr['ok'] and dr['to'] == 'ap@beta.test' and 'HL/002' in dr['subject'] and 'overdue by 40 days' in dr['body']
    assert '₹2.32L' in dr['body']
    assert not a.get('/api/finance/reminder?invoice_id=99731').get_json()['ok']        # already paid
    sent = []
    monkeypatch.setattr(srv, 'email_service_send', lambda to, s, b, body_html=None, **k: (sent.append((to, s)) or (True, '', 'm', 't')))
    assert a.post('/api/finance/reminder', json={'invoice_id': 99732, 'to': 'bad'}).status_code == 400
    assert a.post('/api/finance/reminder', json={'invoice_id': 99732, 'to': 'ap@beta.test', 'subject': dr['subject'], 'body': dr['body']}).get_json()['ok']
    assert sent == [('ap@beta.test', dr['subject'])]
    ag = a.get('/api/finance/summary').get_json()['aging']['invoices']
    assert [x for x in ag if x['id'] == 99732][0]['reminders'] == 1


def test_receipt_and_ai(srv, monkeypatch):
    a = cl(srv, ADMIN)
    monkeypatch.setattr(srv, 'get_setting', lambda k, *x: '')
    r = a.post('/api/finance/receipt', data={'file': (io.BytesIO(b'\xff\xd8\xff fake jpg'), 'bill.jpg')},
               content_type='multipart/form-data').get_json()
    assert r['ok'] and r['fields'] is None and 'Claude' in r['note'] and r['receipt_path'].startswith('r%d_' % CO)
    e = a.post('/api/finance/expenses', json={'amount': 500, 'category': 'Travel', 'receipt_path': r['receipt_path'], 'receipt_name': 'bill.jpg'}).get_json()
    assert a.get('/api/finance/receipt/%d' % e['id']).status_code == 200
    assert cl(srv, OADMIN).get('/api/finance/receipt/%d' % e['id']).status_code == 404
    # someone else's file name is refused
    e2 = a.post('/api/finance/expenses', json={'amount': 5, 'receipt_path': 'r%d_x.jpg' % OTHER}).get_json()
    assert a.get('/api/finance/receipt/%d' % e2['id']).status_code == 404
    assert a.post('/api/finance/receipt', data={'file': (io.BytesIO(b'x'), 'a.exe')}, content_type='multipart/form-data').status_code == 400

    class R:
        status_code = 200
        def __init__(self, t): self.t = t
        def json(self): return {'choices': [{'message': {'content': self.t}}]}
    seen = {}
    monkeypatch.setattr(srv, 'get_setting', lambda k, *x: 'key' if k == 'deepseek_api_key' else '')
    monkeypatch.setattr(srv, 'call_deepseek', lambda key, payload, timeout=60, endpoint='': (seen.update(p=payload) or R('• Revenue ₹3L')))
    j = a.post('/api/finance/insight', json={'period': 'this_month'}).get_json()
    assert j['ok'] and j['text'] == '• Revenue ₹3L' and '300000' in seen['p']['messages'][1]['content']
    assert a.get('/api/finance/summary?period=this_month').get_json()['insight']['text'] == '• Revenue ₹3L'


def test_export_and_page(srv):
    a = cl(srv, ADMIN)
    r = a.get('/api/finance/export?period=fy')
    assert r.status_code == 200 and 'spreadsheetml' in r.mimetype
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ['P&L', 'Invoices', 'Expenses', 'GST', 'TDS']
    nos = [row[0].value for row in wb['Invoices'].iter_rows(min_row=2)]
    assert 'HL/001' in nos and 'PF/001' not in nos
    p = a.get('/finance')
    assert p.status_code == 200 and b'Profit &amp; Cash' in p.data and b'/app-nav.js' in p.data
    assert srv.app.test_client().get('/finance').status_code == 302


def test_periods():
    from modules import finance as f
    ref = datetime.date(2026, 2, 10)
    assert f.period_range('this_quarter', ref=ref)[:2] == (datetime.date(2026, 1, 1), datetime.date(2026, 3, 31))
    assert f.period_range('fy', ref=ref)[:2] == (datetime.date(2025, 4, 1), datetime.date(2026, 3, 31))
    assert f.period_range('last_fy', ref=ref)[2] == 'FY 2024-25'
    assert f.period_range('last_month', ref=datetime.date(2026, 1, 5))[:2] == (datetime.date(2025, 12, 1), datetime.date(2025, 12, 31))
    assert f.period_range('this_quarter', ref=datetime.date(2026, 11, 3))[:2] == (datetime.date(2026, 10, 1), datetime.date(2026, 12, 31))
    assert f.ckey('Larsen & Toubro Ltd.') == f.ckey('LARSEN AND TOUBRO LIMITED'.replace('AND', '&'))

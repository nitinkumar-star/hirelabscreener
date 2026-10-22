"""
Profit & Cash — the rebuilt "Expenses & profit" area of Invoicing (Oct 2026).

What was wrong before (root causes):
  * Net profit = received − expenses, where "received" INCLUDED 18% GST → profit
    was inflated by the GST that belongs to the government.
  * TDS deducted by clients was not tracked anywhere.
  * Only all-time totals; no month / quarter / FY, so nothing to compare with the
    monthly billing target.
  * "Probable" revenue only counted candidates whose offered CTC was filled in,
    which never happens at shared / interview stage → always ₹0.
  * Expenses could not be edited, had no bill, no GST, no link to a client,
    placement or freelancer.

What this module does (company admins only):
  P&L        Revenue = invoice value WITHOUT GST. Direct costs (freelancer
             commission/salary, job boards) → Gross profit. Operating costs →
             Net profit + margin. Expense cost = amount − GST input credit.
  Views      Billing view (by invoice date) or Cash view (money in / money out).
  Periods    This month · Last month · This quarter · This FY · Last FY · Custom.
  Trend      12 months of revenue / costs / profit vs the monthly target.
  Target     ₹ billed this month vs target, gap, placements still needed.
  GST        Per month: GST on invoices − GST on expenses = estimated payable.
  TDS        Per client per FY: expected / deducted (client TDS % or default 10%).
  Unit econ  Profit per client (with payment days), per placement, per person;
             job-board cost per hire.
  Expenses   Add / edit / delete, GST, payment mode, paid/unpaid, link to client /
             job / candidate / freelancer, bill upload with AI reading (PDF text →
             DeepSeek; photo → Claude if its key is set), recurring monthly
             expenses, category budgets with alerts.
  Freelancer When a freelancer-sourced candidate joins (or is invoiced), their
  commission Commission is added as an unpaid expense automatically: 15% of the fee by
             default (Setup), or a per-freelancer fixed ₹ / %. Only placements from
             the 'commission from' date (default: first day finance is opened).
  Cash       90-day forecast (invoices due + committed joinings − recurring −
             unpaid commissions), collections aging 0–30/31–60/61–90/90+, average
             payment days per client, one-click payment reminder e-mail.
  Leakage    Candidates who joined but have no invoice yet.
  Insight    AI summary of the period (DeepSeek, numbers only — never invents).
  Export     One Excel for the CA: P&L, invoices, expenses, GST, TDS.

Everything is an ESTIMATE for management; GST/TDS filing must be checked by a CA.
Additive only: new columns on `expenses`; new tables finance_settings,
finance_client, finance_invoice_meta, finance_recurring, finance_budget,
finance_freelancer_rate, finance_insight.
"""

import datetime
import html
import io
import json
import math
import os
import re

from flask import Blueprint, request, jsonify, Response, redirect, send_file

from modules.shared import get_db, ts, effective_company_id, real_user_id, is_company_admin, login_required, _core
from modules import register_migration

bp = Blueprint('finance', __name__, url_prefix='/api/finance')
pages = Blueprint('finance_pages', __name__)

IST = datetime.timedelta(hours=5, minutes=30)
DIRECT_CATS = ['Freelancer Commission', 'Freelancer Salary', 'Job Boards & Sourcing']
OPERATING_CATS = ['Owner Salary', 'Staff Salary', 'Tools & Subscriptions', 'Marketing', 'Office & Admin', 'Rent',
                  'Travel', 'Bank & Finance Charges', 'Taxes & Fees', 'Other']
CATEGORIES = DIRECT_CATS + OPERATING_CATS
PAY_MODES = ['Bank transfer', 'UPI', 'Card', 'Cash', 'Cheque', 'Other']
DEFAULTS = {'monthly_target': 3000000.0, 'tds_percent': 10.0, 'gst_registered': 1, 'default_fee_percent': 8.33,
            'default_payment_days': 30, 'freelancer_percent': 15.0, 'commission_from': ''}
EXP_COLS = [('gst_amount', 'REAL DEFAULT 0'), ('kind', "TEXT DEFAULT ''"), ('payment_mode', "TEXT DEFAULT ''"),
            ('status', "TEXT DEFAULT 'paid'"), ('client_key', "TEXT DEFAULT ''"), ('client_name', "TEXT DEFAULT ''"),
            ('mandate_id', 'INTEGER DEFAULT 0'), ('candidate_id', 'INTEGER DEFAULT 0'),
            ('freelancer_user_id', 'INTEGER DEFAULT 0'), ('receipt_path', "TEXT DEFAULT ''"),
            ('receipt_name', "TEXT DEFAULT ''"), ('source_key', "TEXT DEFAULT ''"), ('recurring_id', 'INTEGER DEFAULT 0'),
            ('paid_date', "TEXT DEFAULT ''"), ('updated_at', "TEXT DEFAULT ''")]


@register_migration
def _migrate(conn):
    for col, defn in EXP_COLS:
        try:
            conn.execute('ALTER TABLE expenses ADD COLUMN %s %s' % (col, defn))
        except Exception:
            pass
    conn.execute('''CREATE TABLE IF NOT EXISTS finance_settings (
        company_id INTEGER PRIMARY KEY, monthly_target REAL DEFAULT 3000000, tds_percent REAL DEFAULT 10,
        gst_registered INTEGER DEFAULT 1, default_fee_percent REAL DEFAULT 8.33, default_payment_days INTEGER DEFAULT 30,
        updated_by INTEGER DEFAULT 0, updated_at TEXT DEFAULT '')''')
    for col, defn in (('freelancer_percent', 'REAL DEFAULT 15'), ('commission_from', "TEXT DEFAULT ''")):
        try:
            conn.execute('ALTER TABLE finance_settings ADD COLUMN %s %s' % (col, defn))
        except Exception:
            pass
    conn.execute('''CREATE TABLE IF NOT EXISTS finance_client (
        company_id INTEGER NOT NULL, client_key TEXT NOT NULL, client_name TEXT DEFAULT '',
        tds_percent REAL DEFAULT -1, payment_days INTEGER DEFAULT 0, reminder_email TEXT DEFAULT '',
        PRIMARY KEY (company_id, client_key))''')
    conn.execute('''CREATE TABLE IF NOT EXISTS finance_invoice_meta (
        invoice_id INTEGER PRIMARY KEY, company_id INTEGER NOT NULL, tds_amount REAL DEFAULT -1,
        last_reminder_at TEXT DEFAULT '', reminders INTEGER DEFAULT 0)''')
    conn.execute('''CREATE TABLE IF NOT EXISTS finance_recurring (
        id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER NOT NULL, category TEXT DEFAULT '', payee TEXT DEFAULT '',
        amount REAL DEFAULT 0, gst_amount REAL DEFAULT 0, day_of_month INTEGER DEFAULT 1, start_month TEXT DEFAULT '',
        end_month TEXT DEFAULT '', payment_mode TEXT DEFAULT '', note TEXT DEFAULT '', active INTEGER DEFAULT 1,
        created_at TEXT DEFAULT '')''')
    conn.execute('''CREATE TABLE IF NOT EXISTS finance_budget (
        company_id INTEGER NOT NULL, category TEXT NOT NULL, monthly_amount REAL DEFAULT 0,
        PRIMARY KEY (company_id, category))''')
    conn.execute('''CREATE TABLE IF NOT EXISTS finance_freelancer_rate (
        company_id INTEGER NOT NULL, user_id INTEGER NOT NULL, mode TEXT DEFAULT 'percent', value REAL DEFAULT 0,
        PRIMARY KEY (company_id, user_id))''')
    conn.execute('''CREATE TABLE IF NOT EXISTS finance_insight (
        company_id INTEGER NOT NULL, period_key TEXT NOT NULL, day TEXT DEFAULT '', text TEXT DEFAULT '',
        PRIMARY KEY (company_id, period_key))''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_exp_owner_date ON expenses(owner_id, date)')
    try:
        conn.execute('CREATE INDEX IF NOT EXISTS idx_exp_source ON expenses(owner_id, source_key)')
    except Exception:
        pass
    conn.commit()


# ── helpers ────────────────────────────────────────────────────────────────
def now_ist():
    return datetime.datetime.utcnow() + IST


def today():
    return now_ist().date()


def _rows(conn, sql, params=()):
    try:
        return [dict(r) for r in conn.execute(sql, params).fetchall()]
    except Exception:
        return []


def _err(msg, code=400):
    return jsonify({'ok': False, 'error': msg}), code


def _oid():
    return int(effective_company_id() or 0)


def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _int(v):
    try:
        return int(float(v or 0))
    except (TypeError, ValueError):
        return 0


def ckey(name):
    """Normalised client name: 'Larsen & Toubro Ltd.' → 'larsen toubro'."""
    s = re.sub(r'[^a-z0-9& ]', ' ', str(name or '').lower())
    s = re.sub(r'\b(pvt|private|ltd|limited|llp|inc|the|india|co|company)\b', ' ', s)
    return re.sub(r'\s+', ' ', s.replace('&', ' ')).strip()


def pdate(s):
    if not s:
        return None
    try:
        d = _core()._parse_invoice_date(str(s))
        if d:
            return d
    except Exception:
        pass
    try:
        return datetime.date.fromisoformat(str(s)[:10])
    except ValueError:
        return None


def inr(v):
    v = float(v or 0)
    sign = '-' if v < 0 else ''
    v = abs(v)
    if v >= 1e7:
        return '%s₹%.2fCr' % (sign, v / 1e7)
    if v >= 1e5:
        return '%s₹%.2fL' % (sign, v / 1e5)
    return '%s₹%s' % (sign, '{:,}'.format(int(round(v))))


def fy_of(d):
    y = d.year if d.month >= 4 else d.year - 1
    return '%d-%02d' % (y, (y + 1) % 100)


def month_key(d):
    return d.strftime('%Y-%m')


def get_settings(conn, oid):
    r = conn.execute('SELECT * FROM finance_settings WHERE company_id=?', (oid,)).fetchone()
    s = dict(DEFAULTS)
    if r:
        keys = r.keys()
        for k in DEFAULTS:
            if k in keys and r[k] is not None:
                s[k] = r[k]
    if not s.get('commission_from'):
        # automatic freelancer commission starts from the day finance was first opened,
        # so commissions already paid for older placements are not added again.
        s['commission_from'] = today().isoformat()
        try:
            conn.execute('INSERT OR IGNORE INTO finance_settings (company_id) VALUES (?)', (oid,))
            conn.execute("UPDATE finance_settings SET commission_from=? WHERE company_id=? AND COALESCE(commission_from,'')=''",
                         (s['commission_from'], oid))
            conn.commit()
        except Exception:
            pass
    return s


def period_range(kind, frm=None, to=None, ref=None):
    """(start, end, label) inclusive dates for a period name."""
    t = ref or today()
    if kind == 'last_month':
        first = t.replace(day=1)
        end = first - datetime.timedelta(days=1)
        return end.replace(day=1), end, end.strftime('%B %Y')
    if kind == 'this_quarter':
        qs = {4: 4, 5: 4, 6: 4, 7: 7, 8: 7, 9: 7, 10: 10, 11: 10, 12: 10, 1: 1, 2: 1, 3: 1}[t.month]
        start = datetime.date(t.year, qs, 1)
        em = qs + 2
        end = (datetime.date(t.year + (1 if em == 12 else 0), (em % 12) + 1, 1) - datetime.timedelta(days=1))
        qn = {4: 1, 7: 2, 10: 3, 1: 4}[qs]
        return start, end, 'Q%d FY %s' % (qn, fy_of(start))
    if kind in ('fy', 'last_fy'):
        y = t.year if t.month >= 4 else t.year - 1
        if kind == 'last_fy':
            y -= 1
        start, end = datetime.date(y, 4, 1), datetime.date(y + 1, 3, 31)
        return start, end, 'FY %s' % fy_of(start)
    if kind == 'custom':
        a, b = pdate(frm), pdate(to)
        if a and b and a <= b:
            return a, b, '%s – %s' % (a.strftime('%d %b %Y'), b.strftime('%d %b %Y'))
    start = t.replace(day=1)
    end = (start.replace(day=28) + datetime.timedelta(days=4)).replace(day=1) - datetime.timedelta(days=1)
    return start, end, start.strftime('%B %Y')


def _in(d, a, b):
    return d is not None and a <= d <= b


# ══════════════════════════════════════════════════════════════════════════
#  Loading the books
# ══════════════════════════════════════════════════════════════════════════
def client_rules(conn, oid):
    return {r['client_key']: r for r in _rows(conn, 'SELECT * FROM finance_client WHERE company_id=?', (oid,))}


def load_invoices(conn, oid, s=None):
    s = s or get_settings(conn, oid)
    rules = client_rules(conn, oid)
    meta = {r['invoice_id']: r for r in _rows(conn, 'SELECT * FROM finance_invoice_meta WHERE company_id=?', (oid,))}
    out = []
    for r in _rows(conn, "SELECT * FROM invoices WHERE owner_id=? AND (doc_type IS NULL OR doc_type='' OR doc_type='tax')", (oid,)):
        taxable = _num(r.get('amount'))
        rate = _num(r.get('gst_rate')) if r.get('gst_rate') is not None else 18.0
        gst = round(taxable * rate / 100.0, 2) if s['gst_registered'] else 0.0
        gross = round(taxable + gst, 2)
        key = ckey(r.get('buyer_name'))
        rule = rules.get(key)
        tdsp = rule['tds_percent'] if (rule and rule['tds_percent'] is not None and rule['tds_percent'] >= 0) else s['tds_percent']
        m = meta.get(r['id'])
        tds = m['tds_amount'] if (m and m['tds_amount'] is not None and m['tds_amount'] >= 0) else round(taxable * tdsp / 100.0, 2)
        paid = str(r.get('status') or '').lower() == 'paid'
        rec = _num(r.get('received_amount'))
        cash = rec if rec > 0 else (round(gross - tds, 2) if paid else 0.0)
        short = round(gross - tds - rec, 2) if (paid and rec > 0) else 0.0
        idate = pdate(r.get('invoice_date')) or pdate(r.get('created_at'))
        due = pdate(r.get('due_date')) or (idate + datetime.timedelta(days=int((rule['payment_days'] if rule and rule['payment_days'] else 0) or s['default_payment_days'])) if idate else None)
        out.append({'id': r['id'], 'no': r.get('invoice_no') or '', 'date': idate, 'due': due, 'buyer': r.get('buyer_name') or '',
                    'client_key': key, 'taxable': taxable, 'gst': gst, 'gross': gross, 'tds': tds, 'tds_percent': tdsp,
                    'paid': paid, 'received': cash, 'received_date': pdate(r.get('received_date')) or (idate if paid else None),
                    'short': short if short > 1 else 0.0, 'candidate_id': _int(r.get('candidate_id')),
                    'mandate_id': _int(r.get('mandate_id')), 'candidate_name': r.get('candidate_name') or '',
                    'role': r.get('role') or '', 'client_id': _int(r.get('client_id')),
                    'last_reminder_at': (m['last_reminder_at'] if m else '') or '', 'reminders': (m['reminders'] if m else 0) or 0})
    return out


def kind_of(cat, kind=''):
    if kind in ('direct', 'operating'):
        return kind
    return 'direct' if cat in DIRECT_CATS or 'freelancer' in str(cat).lower() or 'naukri' in str(cat).lower() else 'operating'


def load_expenses(conn, oid):
    out = []
    for r in _rows(conn, 'SELECT * FROM expenses WHERE owner_id=? ORDER BY date DESC, id DESC', (oid,)):
        amt = _num(r.get('amount'))
        g = _num(r.get('gst_amount'))
        cat = r.get('category') or 'Other'
        out.append(dict(r, amount=amt, gst_amount=g, net_cost=round(amt - g, 2), d=pdate(r.get('date')) or pdate(r.get('created_at')),
                        kind=kind_of(cat, r.get('kind') or ''), status=(r.get('status') or 'paid'), category=cat))
    return out


# ══════════════════════════════════════════════════════════════════════════
#  Automatic entries: recurring expenses + freelancer commission
# ══════════════════════════════════════════════════════════════════════════
def post_recurring(conn, oid, upto=None):
    upto = upto or today()
    n = 0
    for r in _rows(conn, 'SELECT * FROM finance_recurring WHERE company_id=? AND active=1', (oid,)):
        try:
            y, m = [int(x) for x in str(r['start_month'] or upto.strftime('%Y-%m')).split('-')[:2]]
        except ValueError:
            continue
        end = str(r['end_month'] or '')
        while (y, m) <= (upto.year, upto.month):
            mk = '%04d-%02d' % (y, m)
            if end and mk > end:
                break
            day = min(max(1, int(r['day_of_month'] or 1)), 28)
            d = datetime.date(y, m, day)
            if d <= upto:
                sk = 'rec:%d:%s' % (r['id'], mk)
                if not conn.execute('SELECT 1 FROM expenses WHERE owner_id=? AND source_key=?', (oid, sk)).fetchone():
                    conn.execute('INSERT INTO expenses (owner_id, date, category, payee, amount, note, invoice_id, created_at, gst_amount, '
                                 'kind, payment_mode, status, source_key, recurring_id, paid_date) VALUES (?,?,?,?,?,?,0,?,?,?,?,?,?,?,?)',
                                 (oid, d.isoformat(), r['category'], r['payee'], r['amount'], (r['note'] or '') + ' (recurring)', ts(),
                                  r['gst_amount'] or 0, kind_of(r['category']), r['payment_mode'] or '', 'paid', sk, r['id'], d.isoformat()))
                    n += 1
            m += 1
            if m > 12:
                y, m = y + 1, 1
    if n:
        conn.commit()
    return n


def placement_fee(c, inv_by_cand, s):
    inv = inv_by_cand.get(c['id'])
    if inv:
        return inv['taxable']
    if _num(c.get('placement_fee')) > 0:
        return _num(c['placement_fee'])
    if _num(c.get('offered_ctc')) > 0:
        return _num(c['offered_ctc']) * (_num(c.get('fee_percent')) or s['default_fee_percent']) / 100.0
    return 0.0


def sync_commissions(conn, oid, invoices=None):
    """Add an unpaid 'Freelancer Commission' expense for each freelancer-sourced
    candidate who joined / was invoiced, for freelancers that have a rate."""
    s = get_settings(conn, oid)
    own = {r['user_id']: r for r in _rows(conn, 'SELECT * FROM finance_freelancer_rate WHERE company_id=?', (oid,))}
    rates = {}
    for u in _rows(conn, "SELECT id FROM users WHERE company_id=? AND role='freelancer_sourcer'", (oid,)):
        r = own.get(u['id']) or {'mode': 'percent', 'value': s['freelancer_percent']}
        if _num(r['value']) > 0:
            rates[u['id']] = r
    if not rates:
        return 0
    start = pdate(s['commission_from']) or today()
    invoices = invoices if invoices is not None else load_invoices(conn, oid, s)
    inv_by_cand = {i['candidate_id']: i for i in invoices if i['candidate_id']}
    users = {r['id']: (r['display_name'] or r['username']) for r in _rows(conn, 'SELECT id, username, display_name FROM users WHERE company_id=?', (oid,))}
    cols = {r[1] for r in conn.execute('PRAGMA table_info(candidates)')}
    if 'sourced_by' not in cols:
        return 0
    extra = ', '.join(c for c in ('placement_fee', 'offered_ctc', 'fee_percent', 'joining_date') if c in cols)
    cands = _rows(conn, "SELECT c.id, c.name, c.mandate_id, c.stage, c.sourced_by, m.client%s FROM candidates c "
                        "LEFT JOIN mandates m ON m.id=c.mandate_id WHERE c.owner_id=? AND COALESCE(c.sourced_by,0)<>0"
                        % ((', c.' + extra.replace(', ', ', c.')) if extra else ''), (oid,))
    n = 0
    for c in cands:
        fl = _int(c['sourced_by'])
        if fl not in rates:
            continue
        joined = str(c.get('stage') or '').strip().lower() in ('joined', 'placed') or c['id'] in inv_by_cand
        if not joined:
            continue
        sk = 'flc:%d' % c['id']
        if conn.execute('SELECT 1 FROM expenses WHERE owner_id=? AND source_key=?', (oid, sk)).fetchone():
            continue
        inv0 = inv_by_cand.get(c['id'])
        when = (inv0['date'] if inv0 and inv0['date'] else None) or pdate(c.get('joining_date')) or today()
        if when < start:
            continue
        rt = rates[fl]
        fee = placement_fee(c, inv_by_cand, s)
        amt = rt['value'] if rt['mode'] == 'fixed' else round(fee * rt['value'] / 100.0, 2)
        if amt <= 0:
            continue
        inv = inv_by_cand.get(c['id'])
        d = (inv['date'] if inv and inv['date'] else today()).isoformat()
        conn.execute('INSERT INTO expenses (owner_id, date, category, payee, amount, note, invoice_id, created_at, gst_amount, kind, '
                     'status, client_key, client_name, mandate_id, candidate_id, freelancer_user_id, source_key) '
                     'VALUES (?,?,?,?,?,?,?,?,0,?,?,?,?,?,?,?,?)',
                     (oid, d, 'Freelancer Commission', users.get(fl, 'Freelancer'), amt,
                      'Commission for %s (%s)' % (c['name'] or 'candidate', ('%g%% of fee' % rt['value']) if rt['mode'] != 'fixed' else 'fixed'),
                      inv['id'] if inv else 0, ts(), 'direct', 'unpaid', ckey(c.get('client')), c.get('client') or '',
                      _int(c['mandate_id']), c['id'], fl, sk))
        n += 1
    if n:
        conn.commit()
    return n


# ══════════════════════════════════════════════════════════════════════════
#  The numbers
# ══════════════════════════════════════════════════════════════════════════
def pnl(invoices, expenses, a, b, view='billing'):
    if view == 'cash':
        inflow = sum(i['received'] for i in invoices if i['paid'] and _in(i['received_date'], a, b))
        gst_in = sum(i['gst'] * (i['received'] / max(1.0, i['gross'] - i['tds'])) for i in invoices if i['paid'] and _in(i['received_date'], a, b))
        out_ex = [e for e in expenses if e['status'] == 'paid' and _in(pdate(e.get('paid_date')) or e['d'], a, b)]
        outflow = sum(e['amount'] for e in out_ex)
        return {'view': 'cash', 'money_in': round(inflow), 'money_out': round(outflow), 'net_cash': round(inflow - outflow),
                'gst_inside_money_in': round(gst_in), 'expenses_paid': len(out_ex)}
    inv = [i for i in invoices if _in(i['date'], a, b)]
    ex = [e for e in expenses if _in(e['d'], a, b)]
    revenue = sum(i['taxable'] for i in inv)
    direct = sum(e['net_cost'] for e in ex if e['kind'] == 'direct')
    opex = sum(e['net_cost'] for e in ex if e['kind'] != 'direct')
    by_cat = {}
    for e in ex:
        by_cat[e['category']] = by_cat.get(e['category'], 0) + e['net_cost']
    gross = revenue - direct
    net = gross - opex
    return {'view': 'billing', 'revenue': round(revenue), 'direct_costs': round(direct), 'gross_profit': round(gross),
            'operating_costs': round(opex), 'net_profit': round(net),
            'gross_margin': round(gross * 100.0 / revenue, 1) if revenue else None,
            'net_margin': round(net * 100.0 / revenue, 1) if revenue else None,
            'gst_collected': round(sum(i['gst'] for i in inv)), 'gst_on_expenses': round(sum(e['gst_amount'] for e in ex)),
            'tds': round(sum(i['tds'] for i in inv)), 'invoices': len(inv), 'expenses': len(ex),
            'costs_by_category': {k: round(v) for k, v in sorted(by_cat.items(), key=lambda x: -x[1])},
            'unpaid_expenses': round(sum(e['amount'] for e in ex if e['status'] != 'paid'))}


def trend(invoices, expenses, s, months=12, ref=None):
    t = ref or today()
    y, m = t.year, t.month
    keys = []
    for _ in range(months):
        keys.append((y, m))
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    out = []
    for (y, m) in reversed(keys):
        a = datetime.date(y, m, 1)
        b = (a.replace(day=28) + datetime.timedelta(days=4)).replace(day=1) - datetime.timedelta(days=1)
        p = pnl(invoices, expenses, a, b)
        out.append({'month': a.strftime("%b '%y"), 'key': a.strftime('%Y-%m'), 'revenue': p['revenue'],
                    'costs': p['direct_costs'] + p['operating_costs'], 'profit': p['net_profit'],
                    'cash_in': round(sum(i['received'] for i in invoices if i['paid'] and _in(i['received_date'], a, b)))})
    return {'months': out, 'target': s['monthly_target']}


def target_block(invoices, s, ref=None):
    t = ref or today()
    a, b, label = period_range('this_month', ref=t)
    billed = sum(i['taxable'] for i in invoices if _in(i['date'], a, b))
    yr = [i['taxable'] for i in invoices if i['date'] and i['date'] >= t - datetime.timedelta(days=365) and i['taxable'] > 0]
    avg = (sum(yr) / len(yr)) if yr else 150000.0
    gap = max(0.0, s['monthly_target'] - billed)
    days_left = (b - t).days + 1
    return {'month': label, 'target': round(s['monthly_target']), 'billed': round(billed), 'gap': round(gap),
            'pct': round(billed * 100.0 / s['monthly_target'], 1) if s['monthly_target'] else 0, 'avg_fee': round(avg),
            'placements_needed': int(math.ceil(gap / avg)) if gap and avg else 0, 'days_left': days_left,
            'avg_from': len(yr)}


def gst_months(invoices, expenses, a, b):
    rows = {}
    for i in invoices:
        if _in(i['date'], a, b):
            r = rows.setdefault(month_key(i['date']), {'output': 0.0, 'input': 0.0})
            r['output'] += i['gst']
    for e in expenses:
        if _in(e['d'], a, b) and e['gst_amount']:
            r = rows.setdefault(month_key(e['d']), {'output': 0.0, 'input': 0.0})
            r['input'] += e['gst_amount']
    out = []
    for k in sorted(rows):
        r = rows[k]
        out.append({'month': datetime.date.fromisoformat(k + '-01').strftime('%b %Y'), 'output_gst': round(r['output']),
                    'input_gst': round(r['input']), 'payable': round(max(0.0, r['output'] - r['input'])),
                    'credit_carry': round(max(0.0, r['input'] - r['output']))})
    return out


def tds_register(invoices, fy=None):
    g = {}
    for i in invoices:
        if not i['date']:
            continue
        f = fy_of(i['date'])
        if fy and f != fy:
            continue
        r = g.setdefault((f, i['client_key']), {'fy': f, 'client': i['buyer'], 'client_key': i['client_key'],
                                                'tds_percent': i['tds_percent'], 'invoiced': 0.0, 'deducted': 0.0,
                                                'expected_unpaid': 0.0, 'invoices': 0})
        r['invoiced'] += i['taxable']
        r['invoices'] += 1
        if i['paid']:
            r['deducted'] += i['tds']
        else:
            r['expected_unpaid'] += i['tds']
    return [dict(v, invoiced=round(v['invoiced']), deducted=round(v['deducted']), expected_unpaid=round(v['expected_unpaid']))
            for v in sorted(g.values(), key=lambda v: (v['fy'], -v['invoiced']), reverse=False)]


def by_client(invoices, expenses, a, b):
    g = {}
    for i in invoices:
        k = i['client_key']
        r = g.setdefault(k, {'client_key': k, 'client': i['buyer'], 'revenue': 0.0, 'costs': 0.0, 'invoices': 0,
                             'outstanding': 0.0, 'pay_days': [], 'overdue': 0.0})
        if _in(i['date'], a, b):
            r['revenue'] += i['taxable']
            r['invoices'] += 1
        if i['paid'] and i['date'] and i['received_date']:
            r['pay_days'].append(max(0, (i['received_date'] - i['date']).days))
        if not i['paid']:
            r['outstanding'] += i['gross'] - i['tds']
            if i['due'] and i['due'] < today():
                r['overdue'] += i['gross'] - i['tds']
    cand_client = {i['candidate_id']: i['client_key'] for i in invoices if i['candidate_id']}
    for e in expenses:
        if not _in(e['d'], a, b):
            continue
        k = e.get('client_key') or cand_client.get(_int(e.get('candidate_id')))
        if k and k in g:
            g[k]['costs'] += e['net_cost']
    out = []
    for r in g.values():
        if not (r['revenue'] or r['outstanding']):
            continue
        net = r['revenue'] - r['costs']
        out.append({'client_key': r['client_key'], 'client': r['client'], 'revenue': round(r['revenue']), 'direct_costs': round(r['costs']),
                    'net': round(net), 'margin': round(net * 100.0 / r['revenue'], 1) if r['revenue'] else None,
                    'invoices': r['invoices'], 'outstanding': round(r['outstanding']), 'overdue': round(r['overdue']),
                    'avg_pay_days': round(sum(r['pay_days']) / len(r['pay_days'])) if r['pay_days'] else None})
    return sorted(out, key=lambda r: -r['revenue'])


def by_placement(invoices, expenses, a, b):
    costs = {}
    for e in expenses:
        if _int(e.get('candidate_id')):
            costs[_int(e['candidate_id'])] = costs.get(_int(e['candidate_id']), 0) + e['net_cost']
    out = []
    for i in invoices:
        if not _in(i['date'], a, b):
            continue
        c = costs.get(i['candidate_id'], 0) if i['candidate_id'] else 0
        out.append({'invoice_no': i['no'], 'candidate': i['candidate_name'] or '—', 'role': i['role'], 'client': i['buyer'],
                    'fee': round(i['taxable']), 'costs': round(c), 'net': round(i['taxable'] - c),
                    'margin': round((i['taxable'] - c) * 100.0 / i['taxable'], 1) if i['taxable'] else None,
                    'paid': i['paid']})
    return sorted(out, key=lambda r: -r['fee'])


def by_person(conn, oid, invoices, expenses, a, b):
    cols = {r[1] for r in conn.execute('PRAGMA table_info(candidates)')}
    who = {}
    ids = [i['candidate_id'] for i in invoices if i['candidate_id'] and _in(i['date'], a, b)]
    if ids and ('sourced_by' in cols or 'created_by' in cols):
        sel = ', '.join(c for c in ('sourced_by', 'created_by') if c in cols)
        for r in _rows(conn, 'SELECT id, %s FROM candidates WHERE id IN (%s)' % (sel, ','.join(str(x) for x in ids))):
            who[r['id']] = _int(r.get('sourced_by')) or _int(r.get('created_by'))
    users = {r['id']: (r['display_name'] or r['username'], r['role']) for r in _rows(conn, 'SELECT id, username, display_name, role FROM users WHERE company_id=?', (oid,))}
    g = {}
    for i in invoices:
        if not _in(i['date'], a, b):
            continue
        u = who.get(i['candidate_id'], 0)
        r = g.setdefault(u, {'revenue': 0.0, 'placements': 0, 'costs': 0.0})
        r['revenue'] += i['taxable']
        r['placements'] += 1
    for e in expenses:
        if _in(e['d'], a, b) and _int(e.get('freelancer_user_id')):
            r = g.setdefault(_int(e['freelancer_user_id']), {'revenue': 0.0, 'placements': 0, 'costs': 0.0})
            r['costs'] += e['net_cost']
    out = []
    for u, r in g.items():
        nm, role = users.get(u, ('Not recorded', ''))
        out.append({'user_id': u, 'name': nm, 'freelancer': role == 'freelancer_sourcer', 'revenue': round(r['revenue']),
                    'placements': r['placements'], 'costs': round(r['costs']), 'net': round(r['revenue'] - r['costs'])})
    return sorted(out, key=lambda r: -r['revenue'])


def cost_per_hire(invoices, expenses, a, b):
    boards = sum(e['net_cost'] for e in expenses if _in(e['d'], a, b) and e['category'] == 'Job Boards & Sourcing')
    hires = sum(1 for i in invoices if _in(i['date'], a, b))
    return {'job_board_spend': round(boards), 'placements': hires, 'per_hire': round(boards / hires) if hires else None}


def budgets(conn, oid, expenses, ref=None):
    t = ref or today()
    a, b, _l = period_range('this_month', ref=t)
    spent = {}
    for e in expenses:
        if _in(e['d'], a, b):
            spent[e['category']] = spent.get(e['category'], 0) + e['net_cost']
    out = []
    for r in _rows(conn, 'SELECT category, monthly_amount FROM finance_budget WHERE company_id=? AND monthly_amount>0', (oid,)):
        sp = spent.get(r['category'], 0)
        pct = sp * 100.0 / r['monthly_amount']
        out.append({'category': r['category'], 'budget': round(r['monthly_amount']), 'spent': round(sp), 'pct': round(pct),
                    'state': 'over' if pct > 100 else ('near' if pct >= 80 else 'ok')})
    return sorted(out, key=lambda r: -r['pct'])


def aging(invoices, ref=None):
    t = ref or today()
    buckets = {'0-30': 0.0, '31-60': 0.0, '61-90': 0.0, '90+': 0.0}
    rows = []
    for i in invoices:
        if i['paid']:
            continue
        due_amt = i['gross'] - i['tds']
        base = i['date'] or t
        days = (t - base).days
        k = '0-30' if days <= 30 else '31-60' if days <= 60 else '61-90' if days <= 90 else '90+'
        buckets[k] += due_amt
        rows.append({'id': i['id'], 'no': i['no'], 'client': i['buyer'], 'client_key': i['client_key'], 'amount': round(due_amt),
                     'invoice_date': i['date'].isoformat() if i['date'] else '', 'due_date': i['due'].isoformat() if i['due'] else '',
                     'days': days, 'overdue_days': (t - i['due']).days if (i['due'] and i['due'] < t) else 0, 'bucket': k,
                     'last_reminder_at': i['last_reminder_at'], 'reminders': i['reminders']})
    rows.sort(key=lambda r: (-r['overdue_days'], -r['amount']))
    return {'buckets': {k: round(v) for k, v in buckets.items()}, 'total': round(sum(buckets.values())), 'invoices': rows}


def not_invoiced(conn, oid, invoices, s):
    cols = {r[1] for r in conn.execute('PRAGMA table_info(candidates)')}
    extra = [c for c in ('placement_fee', 'offered_ctc', 'fee_percent', 'joining_date') if c in cols]
    if not extra:
        return []
    have = {i['candidate_id'] for i in invoices if i['candidate_id']}
    names = {(i['candidate_name'] or '').strip().lower() for i in invoices if i['candidate_name']}
    t = today().isoformat()
    out = []
    for c in _rows(conn, "SELECT c.id, c.name, c.stage, c.mandate_id, m.role, m.client, %s FROM candidates c "
                         "LEFT JOIN mandates m ON m.id=c.mandate_id WHERE c.owner_id=? AND (LOWER(TRIM(c.stage)) IN ('joined','placed') "
                         "OR (COALESCE(c.joining_date,'')<>'' AND c.joining_date<=?))" % ', '.join('c.' + x for x in extra), (oid, t)):
        if c['id'] in have or (c['name'] or '').strip().lower() in names:
            continue
        if re.search(r'drop|reject|declin|back ?out|withdr|not join|no show', str(c['stage'] or '').lower()):
            continue
        fee = placement_fee(c, {}, s)
        out.append({'candidate_id': c['id'], 'mandate_id': c['mandate_id'], 'name': c['name'] or '', 'role': c.get('role') or '',
                    'client': c.get('client') or '', 'stage': c['stage'] or '', 'joining_date': c.get('joining_date') or '',
                    'expected_fee': round(fee)})
    return sorted(out, key=lambda r: -r['expected_fee'])


def forecast(conn, oid, invoices, expenses, s, days=90):
    t = today()
    end = t + datetime.timedelta(days=days)
    weeks = []
    for w in range(0, days, 7):
        weeks.append({'from': (t + datetime.timedelta(days=w)).isoformat(), 'in': 0.0, 'out': 0.0})

    def put(d, k, v):
        if d is None:
            d = t
        idx = max(0, (d - t).days) // 7
        if idx < len(weeks):
            weeks[idx][k] += v
    rules = client_rules(conn, oid)
    inflow = []
    for i in invoices:
        if i['paid']:
            continue
        d = i['due'] if (i['due'] and i['due'] >= t) else t
        if d <= end:
            put(d, 'in', i['gross'] - i['tds'])
            inflow.append({'what': 'Invoice %s · %s' % (i['no'], i['buyer']), 'when': d.isoformat(), 'amount': round(i['gross'] - i['tds']),
                           'overdue': bool(i['due'] and i['due'] < t)})
    # committed joinings (future joining date, not invoiced) → paid about N days after joining
    cols = {r[1] for r in conn.execute('PRAGMA table_info(candidates)')}
    if 'joining_date' in cols:
        have = {i['candidate_id'] for i in invoices if i['candidate_id']}
        for c in _rows(conn, "SELECT c.id, c.name, c.joining_date, c.placement_fee, c.offered_ctc, c.fee_percent, m.client FROM candidates c "
                             "LEFT JOIN mandates m ON m.id=c.mandate_id WHERE c.owner_id=? AND COALESCE(c.joining_date,'')>? "
                             "AND (COALESCE(c.offered_ctc,0)>0 OR COALESCE(c.placement_fee,0)>0)", (oid, t.isoformat())):
            if c['id'] in have:
                continue
            jd = pdate(c['joining_date'])
            if not jd:
                continue
            rule = rules.get(ckey(c.get('client')))
            pdays = (rule['payment_days'] if rule and rule['payment_days'] else 0) or s['default_payment_days']
            d = jd + datetime.timedelta(days=int(pdays))
            fee = placement_fee(c, {}, s)
            tdsp = rule['tds_percent'] if (rule and rule['tds_percent'] is not None and rule['tds_percent'] >= 0) else s['tds_percent']
            amt = fee * (1 + (0.18 if s['gst_registered'] else 0)) - fee * tdsp / 100.0
            if d <= end and amt > 0:
                put(d, 'in', amt)
                inflow.append({'what': 'Joining: %s · %s' % (c['name'], c.get('client') or ''), 'when': d.isoformat(),
                               'amount': round(amt), 'overdue': False})
    outflow = []
    for r in _rows(conn, 'SELECT * FROM finance_recurring WHERE company_id=? AND active=1', (oid,)):
        d = t.replace(day=min(28, max(1, int(r['day_of_month'] or 1))))
        if d < t:
            d = (d.replace(day=1) + datetime.timedelta(days=32)).replace(day=d.day)
        while d <= end:
            if not r['end_month'] or d.strftime('%Y-%m') <= r['end_month']:
                put(d, 'out', r['amount'])
                outflow.append({'what': '%s · %s' % (r['category'], r['payee']), 'when': d.isoformat(), 'amount': round(r['amount'])})
            d = (d.replace(day=1) + datetime.timedelta(days=32)).replace(day=d.day)
    for e in expenses:
        if e['status'] != 'paid':
            put(t, 'out', e['amount'])
            outflow.append({'what': 'Unpaid: %s · %s' % (e['category'], e.get('payee') or ''), 'when': t.isoformat(), 'amount': round(e['amount'])})
    run = 0.0
    for w in weeks:
        run += w['in'] - w['out']
        w['in'], w['out'], w['net_running'] = round(w['in']), round(w['out']), round(run)
    return {'weeks': weeks, 'total_in': round(sum(w['in'] for w in weeks)), 'total_out': round(sum(w['out'] for w in weeks)),
            'inflows': sorted(inflow, key=lambda x: x['when'])[:40], 'outflows': sorted(outflow, key=lambda x: x['when'])[:40]}


# ══════════════════════════════════════════════════════════════════════════
#  API
# ══════════════════════════════════════════════════════════════════════════
def _admin():
    if not is_company_admin():
        return _err('Only an admin can see finances.', 403)
    return None


def _period_from_args():
    return period_range(request.args.get('period') or 'this_month', request.args.get('from'), request.args.get('to'))


@bp.route('/summary', methods=['GET'])
@login_required
def summary():
    e = _admin()
    if e:
        return e
    oid = _oid()
    conn = get_db()
    try:
        s = get_settings(conn, oid)
        post_recurring(conn, oid)
        invoices = load_invoices(conn, oid, s)
        sync_commissions(conn, oid, invoices)
        expenses = load_expenses(conn, oid)
        a, b, label = _period_from_args()
        view = 'cash' if request.args.get('view') == 'cash' else 'billing'
        p = pnl(invoices, expenses, a, b)
        out = {'ok': True, 'period': {'from': a.isoformat(), 'to': b.isoformat(), 'label': label, 'key': request.args.get('period') or 'this_month'},
               'settings': s, 'pnl': p, 'cash': pnl(invoices, expenses, a, b, 'cash'), 'view': view,
               'trend': trend(invoices, expenses, s), 'target': target_block(invoices, s),
               'gst': gst_months(invoices, expenses, a, b), 'tds': tds_register(invoices, fy_of(b)),
               'clients': by_client(invoices, expenses, a, b), 'placements': by_placement(invoices, expenses, a, b),
               'people': by_person(conn, oid, invoices, expenses, a, b), 'cost_per_hire': cost_per_hire(invoices, expenses, a, b),
               'budgets': budgets(conn, oid, expenses), 'aging': aging(invoices), 'not_invoiced': not_invoiced(conn, oid, invoices, s),
               'forecast': forecast(conn, oid, invoices, expenses, s),
               'short_paid': [{'no': i['no'], 'client': i['buyer'], 'short': round(i['short'])} for i in invoices if i['short']][:20],
               'categories': CATEGORIES, 'direct_categories': DIRECT_CATS, 'pay_modes': PAY_MODES}
        ins = conn.execute('SELECT text, day FROM finance_insight WHERE company_id=? AND period_key=?',
                           (oid, '%s:%s' % (a.isoformat(), b.isoformat()))).fetchone()
        out['insight'] = {'text': ins['text'], 'day': ins['day']} if ins else None
        return jsonify(out)
    finally:
        conn.close()


@bp.route('/expenses', methods=['GET'])
@login_required
def expenses_list():
    e = _admin()
    if e:
        return e
    oid = _oid()
    conn = get_db()
    try:
        a, b, label = _period_from_args()
        rows = [x for x in load_expenses(conn, oid) if request.args.get('all') or _in(x['d'], a, b)]
        cat = request.args.get('category') or ''
        if cat:
            rows = [x for x in rows if x['category'] == cat]
        for x in rows:
            x.pop('d', None)
            x['has_receipt'] = bool(x.get('receipt_path'))
            x.pop('receipt_path', None)
        return jsonify({'ok': True, 'expenses': rows, 'period': label})
    finally:
        conn.close()


def _expense_values(d):
    cat = str(d.get('category') or 'Other').strip()[:80]
    amt = _num(d.get('amount'))
    if amt <= 0:
        return None, 'Amount must be more than 0.'
    g = _num(d.get('gst_amount'))
    if g < 0 or g >= amt:
        return None, 'GST must be less than the amount (the amount includes GST).'
    dt = pdate(d.get('date')) or today()
    status = 'unpaid' if d.get('status') == 'unpaid' else 'paid'
    client = str(d.get('client_name') or '').strip()[:120]
    return {'date': dt.isoformat(), 'category': cat, 'payee': str(d.get('payee') or '').strip()[:120], 'amount': amt,
            'note': str(d.get('note') or '').strip()[:500], 'gst_amount': g, 'kind': kind_of(cat, d.get('kind') or ''),
            'payment_mode': str(d.get('payment_mode') or '')[:30], 'status': status,
            'paid_date': (pdate(d.get('paid_date')) or dt).isoformat() if status == 'paid' else '',
            'client_name': client, 'client_key': ckey(client) if client else '',
            'mandate_id': _int(d.get('mandate_id')), 'candidate_id': _int(d.get('candidate_id')),
            'freelancer_user_id': _int(d.get('freelancer_user_id')), 'updated_at': ts()}, None


@bp.route('/expenses', methods=['POST'])
@login_required
def expense_create():
    e = _admin()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    v, er = _expense_values(d)
    if er:
        return _err(er)
    oid = _oid()
    conn = get_db()
    try:
        v.update(owner_id=oid, created_at=ts(), invoice_id=0, receipt_path=str(d.get('receipt_path') or '')[:200],
                 receipt_name=str(d.get('receipt_name') or '')[:200])
        if v['receipt_path'] and not re.match(r'^r%d_[A-Za-z0-9_.-]+$' % oid, v['receipt_path']):
            v['receipt_path'] = ''
        cur = conn.execute('INSERT INTO expenses (%s) VALUES (%s)' % (','.join(v), ','.join('?' * len(v))), list(v.values()))
        if d.get('recurring'):
            mk = v['date'][:7]
            conn.execute('INSERT INTO finance_recurring (company_id, category, payee, amount, gst_amount, day_of_month, start_month, '
                         'payment_mode, note, active, created_at) VALUES (?,?,?,?,?,?,?,?,?,1,?)',
                         (oid, v['category'], v['payee'], v['amount'], v['gst_amount'], int(v['date'][8:10]),
                          mk, v['payment_mode'], v['note'], ts()))
            rid = conn.execute('SELECT last_insert_rowid()').fetchone()[0]
            conn.execute('UPDATE expenses SET recurring_id=?, source_key=? WHERE id=?', (rid, 'rec:%d:%s' % (rid, mk), cur.lastrowid))
        conn.commit()
        return jsonify({'ok': True, 'id': cur.lastrowid})
    finally:
        conn.close()


@bp.route('/expense/<int:eid>', methods=['PUT', 'POST'])
@login_required
def expense_update(eid):
    e = _admin()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    oid = _oid()
    conn = get_db()
    try:
        cur = conn.execute('SELECT * FROM expenses WHERE id=? AND owner_id=?', (eid, oid)).fetchone()
        if not cur:
            return _err('Not found', 404)
        if set(d) <= {'status', 'paid_date'}:          # quick "mark paid / unpaid"
            st = 'unpaid' if d.get('status') == 'unpaid' else 'paid'
            conn.execute('UPDATE expenses SET status=?, paid_date=?, updated_at=? WHERE id=?',
                         (st, ((pdate(d.get('paid_date')) or today()).isoformat() if st == 'paid' else ''), ts(), eid))
        else:
            v, er = _expense_values(dict(dict(cur), **d))
            if er:
                return _err(er)
            conn.execute('UPDATE expenses SET %s WHERE id=?' % ','.join(k + '=?' for k in v), list(v.values()) + [eid])
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


@bp.route('/expense/<int:eid>', methods=['DELETE'])
@login_required
def expense_delete(eid):
    e = _admin()
    if e:
        return e
    oid = _oid()
    conn = get_db()
    try:
        r = conn.execute('SELECT source_key FROM expenses WHERE id=? AND owner_id=?', (eid, oid)).fetchone()
        if not r:
            return _err('Not found', 404)
        if str(r['source_key'] or '').startswith('flc:'):
            # keep a tombstone so the automatic commission is not re-created
            conn.execute("UPDATE expenses SET amount=0, gst_amount=0, note='(commission removed)', status='paid', category='Freelancer Commission' WHERE id=?", (eid,))
        else:
            conn.execute('DELETE FROM expenses WHERE id=?', (eid,))
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


# ── receipts + AI reading ─────────────────────────────────────────────────
def _receipt_dir():
    d = os.path.join(_core().DATA_DIR, 'receipts')
    os.makedirs(d, exist_ok=True)
    return d


_BILL_PROMPT = ('Read this bill / receipt / invoice for an Indian business expense. Reply ONLY with JSON: '
                '{"date":"YYYY-MM-DD or empty","payee":"seller / vendor name","amount":total amount paid including GST as a number,'
                '"gst_amount":total GST (CGST+SGST+IGST) as a number or 0,"category":one of %s,'
                '"note":"what was bought, max 8 words"}. Use only what is written on the bill.' % json.dumps(CATEGORIES))


def _read_bill(path, name):
    core = _core()
    ext = os.path.splitext(name.lower())[1]
    if ext == '.pdf':
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                text = '\n'.join((p.extract_text() or '') for p in pdf.pages[:3])[:6000]
        except Exception as exc:
            return None, 'Could not read the PDF (%s)' % str(exc)[:60]
        if not text.strip():
            return None, 'This PDF is a scanned image; fill the fields by hand (or upload a photo with the Claude key set).'
        key = (core.get_setting('deepseek_api_key') or '').strip()
        if not key:
            return None, 'DeepSeek key not set — fill the fields by hand.'
        resp = core.call_deepseek(key, {'model': 'deepseek-chat', 'temperature': 0, 'max_tokens': 400,
                                        'response_format': {'type': 'json_object'},
                                        'messages': [{'role': 'system', 'content': _BILL_PROMPT}, {'role': 'user', 'content': text}]},
                                  timeout=60, endpoint='finance-bill')
        raw = resp.json()['choices'][0]['message']['content'] if resp.status_code == 200 else ''
    elif ext in ('.jpg', '.jpeg', '.png', '.webp'):
        key = (core.get_setting('claude_api_key') or '').strip()
        if not key:
            return None, 'Reading photos needs the Claude API key (Settings). Fill the fields by hand, or upload the PDF bill.'
        import base64
        with open(path, 'rb') as f:
            data = base64.b64encode(f.read()).decode()
        mt = 'image/png' if ext == '.png' else 'image/webp' if ext == '.webp' else 'image/jpeg'
        resp = core.call_claude(key, _BILL_PROMPT, [{'role': 'user', 'content': [
            {'type': 'image', 'source': {'type': 'base64', 'media_type': mt, 'data': data}},
            {'type': 'text', 'text': 'Read this bill.'}]}], max_tokens=500, endpoint='finance-bill')
        raw = ''
        if resp.status_code == 200:
            raw = ''.join(b.get('text', '') for b in resp.json().get('content', []) if b.get('type') == 'text')
    else:
        return None, 'Upload a PDF, JPG or PNG.'
    try:
        s, e = raw.find('{'), raw.rfind('}')
        j = json.loads(raw[s:e + 1])
    except Exception:
        return None, 'Could not read this bill automatically — fill the fields by hand.'
    cat = j.get('category') if j.get('category') in CATEGORIES else 'Other'
    return {'date': str(j.get('date') or '')[:10], 'payee': str(j.get('payee') or '')[:120], 'amount': _num(j.get('amount')),
            'gst_amount': _num(j.get('gst_amount')), 'category': cat, 'note': str(j.get('note') or '')[:120]}, None


@bp.route('/receipt', methods=['POST'])
@login_required
def receipt_upload():
    e = _admin()
    if e:
        return e
    f = request.files.get('file')
    if not f or not f.filename:
        return _err('Choose a file.')
    ext = os.path.splitext(f.filename.lower())[1]
    if ext not in ('.pdf', '.jpg', '.jpeg', '.png', '.webp'):
        return _err('Upload a PDF, JPG or PNG.')
    oid = _oid()
    name = 'r%d_%s%s' % (oid, now_ist().strftime('%Y%m%d%H%M%S%f'), ext)
    path = os.path.join(_receipt_dir(), name)
    f.save(path)
    if os.path.getsize(path) > 10 * 1024 * 1024:
        os.remove(path)
        return _err('Max 10 MB.')
    fields, why = None, ''
    try:
        fields, why = _read_bill(path, f.filename)
    except Exception as exc:
        why = 'Could not read automatically (%s) — fill the fields by hand.' % str(exc)[:80]
    return jsonify({'ok': True, 'receipt_path': name, 'receipt_name': f.filename[:200], 'fields': fields, 'note': why or ''})


@bp.route('/receipt/<int:eid>', methods=['GET'])
@login_required
def receipt_get(eid):
    e = _admin()
    if e:
        return e
    oid = _oid()
    conn = get_db()
    try:
        r = conn.execute('SELECT receipt_path, receipt_name FROM expenses WHERE id=? AND owner_id=?', (eid, oid)).fetchone()
    finally:
        conn.close()
    if not r or not r['receipt_path'] or not re.match(r'^r%d_[A-Za-z0-9_.-]+$' % oid, r['receipt_path']):
        return _err('Not found', 404)
    p = os.path.join(_receipt_dir(), r['receipt_path'])
    if not os.path.isfile(p):
        return _err('Not found', 404)
    return send_file(p, download_name=r['receipt_name'] or r['receipt_path'])


# ── settings / clients / recurring / budgets / freelancer rates ───────────
@bp.route('/setup', methods=['GET'])
@login_required
def setup_get():
    e = _admin()
    if e:
        return e
    oid = _oid()
    conn = get_db()
    try:
        names = {}
        for r in _rows(conn, "SELECT DISTINCT buyer_name FROM invoices WHERE owner_id=? AND COALESCE(buyer_name,'')<>''", (oid,)):
            names.setdefault(ckey(r['buyer_name']), r['buyer_name'])
        rules = client_rules(conn, oid)
        clients = [{'client_key': k, 'client': names.get(k) or (rules[k]['client_name'] if k in rules else k),
                    'tds_percent': rules[k]['tds_percent'] if k in rules and rules[k]['tds_percent'] is not None and rules[k]['tds_percent'] >= 0 else None,
                    'payment_days': rules[k]['payment_days'] if k in rules else 0,
                    'reminder_email': rules[k]['reminder_email'] if k in rules else ''} for k in sorted(set(names) | set(rules))]
        fl = _rows(conn, "SELECT u.id, COALESCE(u.display_name, u.username) AS name, r.mode, r.value FROM users u "
                         "LEFT JOIN finance_freelancer_rate r ON r.user_id=u.id AND r.company_id=u.company_id "
                         "WHERE u.company_id=? AND u.role='freelancer_sourcer'", (oid,))
        return jsonify({'ok': True, 'settings': get_settings(conn, oid), 'clients': clients,
                        'recurring': _rows(conn, 'SELECT * FROM finance_recurring WHERE company_id=? ORDER BY active DESC, id DESC', (oid,)),
                        'budgets': _rows(conn, 'SELECT category, monthly_amount FROM finance_budget WHERE company_id=?', (oid,)),
                        'freelancers': fl, 'categories': CATEGORIES, 'pay_modes': PAY_MODES})
    finally:
        conn.close()


@bp.route('/setup', methods=['PUT', 'POST'])
@login_required
def setup_save():
    e = _admin()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    oid = _oid()
    conn = get_db()
    try:
        if isinstance(d.get('settings'), dict):
            s = get_settings(conn, oid)
            x = d['settings']
            vals = {'monthly_target': max(0.0, _num(x.get('monthly_target', s['monthly_target']))),
                    'tds_percent': min(30.0, max(0.0, _num(x.get('tds_percent', s['tds_percent'])))),
                    'gst_registered': 1 if x.get('gst_registered', s['gst_registered']) in (1, True, '1', 'true') else 0,
                    'default_fee_percent': min(50.0, max(0.0, _num(x.get('default_fee_percent', s['default_fee_percent'])))),
                    'default_payment_days': min(180, max(0, _int(x.get('default_payment_days', s['default_payment_days'])))),
                    'freelancer_percent': min(100.0, max(0.0, _num(x.get('freelancer_percent', s['freelancer_percent'])))),
                    'commission_from': (pdate(x.get('commission_from')) or pdate(s['commission_from']) or today()).isoformat()}
            conn.execute('INSERT OR IGNORE INTO finance_settings (company_id) VALUES (?)', (oid,))
            conn.execute('UPDATE finance_settings SET %s, updated_by=?, updated_at=? WHERE company_id=?' % ','.join(k + '=?' for k in vals),
                         list(vals.values()) + [int(real_user_id() or 0), ts(), oid])
        for c in (d.get('clients') or [])[:500]:
            k = ckey(c.get('client') or c.get('client_key'))
            if not k:
                continue
            tdsp = c.get('tds_percent')
            tdsp = -1 if tdsp in (None, '') else min(30.0, max(0.0, _num(tdsp)))
            conn.execute('INSERT INTO finance_client (company_id, client_key, client_name, tds_percent, payment_days, reminder_email) '
                         'VALUES (?,?,?,?,?,?) ON CONFLICT(company_id, client_key) DO UPDATE SET client_name=excluded.client_name, '
                         'tds_percent=excluded.tds_percent, payment_days=excluded.payment_days, reminder_email=excluded.reminder_email',
                         (oid, k, str(c.get('client') or '')[:120], tdsp, min(180, max(0, _int(c.get('payment_days')))),
                          str(c.get('reminder_email') or '').strip()[:200]))
        for r in (d.get('recurring') or [])[:100]:
            if r.get('delete') and _int(r.get('id')):
                conn.execute('UPDATE finance_recurring SET active=0 WHERE id=? AND company_id=?', (_int(r['id']), oid))
                continue
            amt = _num(r.get('amount'))
            if amt <= 0 or not r.get('category'):
                continue
            vals = (str(r['category'])[:80], str(r.get('payee') or '')[:120], amt, max(0.0, _num(r.get('gst_amount'))),
                    min(28, max(1, _int(r.get('day_of_month')) or 1)), str(r.get('start_month') or today().strftime('%Y-%m'))[:7],
                    str(r.get('end_month') or '')[:7], str(r.get('payment_mode') or '')[:30], str(r.get('note') or '')[:200])
            if _int(r.get('id')):
                conn.execute('UPDATE finance_recurring SET category=?, payee=?, amount=?, gst_amount=?, day_of_month=?, start_month=?, '
                             'end_month=?, payment_mode=?, note=? WHERE id=? AND company_id=?', vals + (_int(r['id']), oid))
            else:
                conn.execute('INSERT INTO finance_recurring (category, payee, amount, gst_amount, day_of_month, start_month, end_month, '
                             'payment_mode, note, company_id, active, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,1,?)', vals + (oid, ts()))
        for bgt in (d.get('budgets') or [])[:50]:
            cat = str(bgt.get('category') or '')[:80]
            if cat:
                conn.execute('INSERT INTO finance_budget (company_id, category, monthly_amount) VALUES (?,?,?) ON CONFLICT(company_id, category) '
                             'DO UPDATE SET monthly_amount=excluded.monthly_amount', (oid, cat, max(0.0, _num(bgt.get('monthly_amount')))))
        for fl in (d.get('freelancers') or [])[:200]:
            uid = _int(fl.get('id'))
            if not conn.execute("SELECT 1 FROM users WHERE id=? AND company_id=?", (uid, oid)).fetchone():
                continue
            if fl.get('value') in (None, ''):          # blank → use the default % from settings
                conn.execute('DELETE FROM finance_freelancer_rate WHERE company_id=? AND user_id=?', (oid, uid))
                continue
            mode = 'fixed' if fl.get('mode') == 'fixed' else 'percent'
            val = max(0.0, _num(fl.get('value')))
            if mode == 'percent':
                val = min(100.0, val)
            conn.execute('INSERT INTO finance_freelancer_rate (company_id, user_id, mode, value) VALUES (?,?,?,?) '
                         'ON CONFLICT(company_id, user_id) DO UPDATE SET mode=excluded.mode, value=excluded.value', (oid, uid, mode, val))
        if isinstance(d.get('invoice_tds'), dict):
            it = d['invoice_tds']
            iid = _int(it.get('invoice_id'))
            if conn.execute('SELECT 1 FROM invoices WHERE id=? AND owner_id=?', (iid, oid)).fetchone():
                v = it.get('tds_amount')
                conn.execute('INSERT INTO finance_invoice_meta (invoice_id, company_id, tds_amount) VALUES (?,?,?) '
                             'ON CONFLICT(invoice_id) DO UPDATE SET tds_amount=excluded.tds_amount',
                             (iid, oid, -1 if v in (None, '') else max(0.0, _num(v))))
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


# ── payment reminder ──────────────────────────────────────────────────────
def _reminder_draft(conn, oid, iid):
    s = get_settings(conn, oid)
    inv = next((i for i in load_invoices(conn, oid, s) if i['id'] == iid), None)
    if not inv:
        return None, 'Invoice not found'
    if inv['paid']:
        return None, 'This invoice is already paid.'
    rule = client_rules(conn, oid).get(inv['client_key'])
    to = (rule['reminder_email'] if rule else '') or ''
    if not to and inv['client_id']:
        r = conn.execute("SELECT email FROM crm_contacts WHERE client_id=? AND company_id=? AND COALESCE(email,'')<>'' AND is_active=1 "
                         "ORDER BY is_primary DESC, id LIMIT 1", (inv['client_id'], oid)).fetchone()
        to = r['email'] if r else ''
    od = (today() - inv['due']).days if inv['due'] and inv['due'] < today() else 0
    body = ('Dear Sir/Madam,\n\nThis is a gentle reminder that invoice %s dated %s for %s is %s. Amount due: %s '
            '(invoice value %s incl. GST, less TDS %s).\n\nRequest you to kindly process the payment at the earliest. '
            'If the payment has already been made, please share the UTR / transaction reference so we can update our records.\n\n'
            'Thank you.' % (inv['no'], inv['date'].strftime('%d %b %Y') if inv['date'] else '-',
                            ('the placement of %s (%s)' % (inv['candidate_name'], inv['role'])) if inv['candidate_name'] else 'our recruitment services',
                            ('overdue by %d days' % od) if od else ('due on %s' % inv['due'].strftime('%d %b %Y') if inv['due'] else 'pending'),
                            inr(inv['gross'] - inv['tds']), inr(inv['gross']), inr(inv['tds'])))
    return {'to': to, 'subject': 'Payment reminder — Invoice %s' % inv['no'], 'body': body, 'client': inv['buyer'],
            'reminders': inv['reminders'], 'last_reminder_at': inv['last_reminder_at']}, None


@bp.route('/reminder', methods=['GET'])
@login_required
def reminder_get():
    e = _admin()
    if e:
        return e
    conn = get_db()
    try:
        d, er = _reminder_draft(conn, _oid(), _int(request.args.get('invoice_id')))
        return _err(er) if er else jsonify(dict(d, ok=True))
    finally:
        conn.close()


@bp.route('/reminder', methods=['POST'])
@login_required
def reminder_send():
    e = _admin()
    if e:
        return e
    core = _core()
    d = request.get_json(silent=True) or {}
    oid = _oid()
    iid = _int(d.get('invoice_id'))
    to = str(d.get('to') or '').strip()
    if not re.match(r'^[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+$', to):
        return _err('Enter the client\'s accounts e-mail.')
    subject, body = str(d.get('subject') or '').strip()[:200], str(d.get('body') or '').strip()[:5000]
    if not subject or not body:
        return _err('Subject and message are needed.')
    conn = get_db()
    try:
        inv = conn.execute('SELECT id, buyer_name FROM invoices WHERE id=? AND owner_id=?', (iid, oid)).fetchone()
        if not inv:
            return _err('Not found', 404)
        bh = ''.join('<p>%s</p>' % html.escape(p).replace('\n', '<br>') for p in re.split(r'\n\s*\n', body) if p.strip())
        ok, err, _m, _t = core.email_service_send(to, subject, body, body_html=bh)
        if not ok:
            return _err('Could not send: %s' % (err or 'mail error'))
        conn.execute('INSERT INTO finance_invoice_meta (invoice_id, company_id, last_reminder_at, reminders) VALUES (?,?,?,1) '
                     'ON CONFLICT(invoice_id) DO UPDATE SET last_reminder_at=excluded.last_reminder_at, reminders=reminders+1',
                     (iid, oid, ts()))
        k = ckey(inv['buyer_name'])
        conn.execute('INSERT INTO finance_client (company_id, client_key, client_name, reminder_email) VALUES (?,?,?,?) '
                     'ON CONFLICT(company_id, client_key) DO UPDATE SET reminder_email=excluded.reminder_email', (oid, k, inv['buyer_name'], to))
        conn.commit()
        return jsonify({'ok': True})
    finally:
        conn.close()


# ── AI insight ────────────────────────────────────────────────────────────
@bp.route('/insight', methods=['POST'])
@login_required
def insight():
    e = _admin()
    if e:
        return e
    core = _core()
    oid = _oid()
    d = request.get_json(silent=True) or {}
    a, b, label = period_range(d.get('period') or 'this_month', d.get('from'), d.get('to'))
    key = (core.get_setting('deepseek_api_key') or '').strip()
    if not key:
        return _err('DeepSeek API key is not set (Settings).')
    conn = get_db()
    try:
        s = get_settings(conn, oid)
        inv = load_invoices(conn, oid, s)
        ex = load_expenses(conn, oid)
        span = (b - a).days + 1
        pa, pb = a - datetime.timedelta(days=span), a - datetime.timedelta(days=1)
        facts = {'period': label, 'this_period': pnl(inv, ex, a, b), 'previous_period_same_length': pnl(inv, ex, pa, pb),
                 'monthly_target': s['monthly_target'], 'target_this_month': target_block(inv, s),
                 'top_clients': by_client(inv, ex, a, b)[:6], 'collections': aging(inv)['buckets'],
                 'overdue_invoices': [x for x in aging(inv)['invoices'] if x['overdue_days']][:6],
                 'joined_not_invoiced': not_invoiced(conn, oid, inv, s)[:6], 'budgets': budgets(conn, oid, ex)}
        system = ('You are the finance assistant of HireLab, an Indian recruitment agency. Write a short management summary '
                  '(5–7 bullet points, plain text, each starting with "• ") of the numbers given. Use ONLY these numbers; never invent. '
                  'Amounts in Indian style (₹ lakh). Cover: revenue vs previous period and vs target, margin and what drove costs, '
                  'clients who pay slowly / overdue money, joined-but-not-invoiced leakage, one or two concrete actions. '
                  'Revenue is without GST. Be direct, no fluff.')
        resp = core.call_deepseek(key, {'model': 'deepseek-chat', 'temperature': 0.2, 'max_tokens': 700,
                                        'messages': [{'role': 'system', 'content': system},
                                                     {'role': 'user', 'content': json.dumps(facts, default=str)[:12000]}]},
                                  timeout=60, endpoint='finance-insight')
        if resp.status_code != 200:
            return _err('AI summary failed (HTTP %s).' % resp.status_code)
        text = (resp.json()['choices'][0]['message']['content'] or '').strip()
        conn.execute('INSERT INTO finance_insight (company_id, period_key, day, text) VALUES (?,?,?,?) ON CONFLICT(company_id, period_key) '
                     'DO UPDATE SET day=excluded.day, text=excluded.text', (oid, '%s:%s' % (a.isoformat(), b.isoformat()), ts()[:10], text))
        conn.commit()
        return jsonify({'ok': True, 'text': text})
    finally:
        conn.close()


# ── Excel for the CA ──────────────────────────────────────────────────────
@bp.route('/export', methods=['GET'])
@login_required
def export_xlsx():
    e = _admin()
    if e:
        return e
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    oid = _oid()
    conn = get_db()
    try:
        s = get_settings(conn, oid)
        inv = load_invoices(conn, oid, s)
        ex = load_expenses(conn, oid)
        a, b, label = _period_from_args()
        p = pnl(inv, ex, a, b)
        wb = Workbook()
        bold, head = Font(bold=True), PatternFill('solid', fgColor='E6F4EE')

        def sheet(ws, header, rows, money_cols=()):
            ws.append(header)
            for c in ws[1]:
                c.font, c.fill = bold, head
            for r in rows:
                ws.append(r)
            for col in money_cols:
                for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    for c in row:
                        c.number_format = '#,##0'
            for i, h in enumerate(header, 1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(42, len(str(h)) + 6))
        ws = wb.active
        ws.title = 'P&L'
        ws.append(['HireLab — Profit & Loss (management estimate)', label])
        ws['A1'].font = Font(bold=True, size=13)
        ws.append([])
        for k, lbl in (('revenue', 'Revenue (invoices, without GST)'), ('direct_costs', 'Direct costs'), ('gross_profit', 'Gross profit'),
                       ('operating_costs', 'Operating costs'), ('net_profit', 'Net profit'), ('gst_collected', 'GST on invoices (output)'),
                       ('gst_on_expenses', 'GST on expenses (input)'), ('tds', 'TDS by clients (expected)')):
            ws.append([lbl, p[k]])
            ws.cell(row=ws.max_row, column=2).number_format = '#,##0'
        ws.append(['Net margin %', p['net_margin']])
        ws.append([])
        ws.append(['Costs by category'])
        ws.cell(row=ws.max_row, column=1).font = bold
        for k, v in p['costs_by_category'].items():
            ws.append([k, v])
            ws.cell(row=ws.max_row, column=2).number_format = '#,##0'
        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 18
        sheet(wb.create_sheet('Invoices'), ['Invoice no', 'Date', 'Client', 'Candidate', 'Role', 'Taxable value', 'GST', 'Invoice total',
                                            'TDS', 'Status', 'Received', 'Received on', 'Due date'],
              [[i['no'], i['date'].isoformat() if i['date'] else '', i['buyer'], i['candidate_name'], i['role'], i['taxable'], i['gst'],
                i['gross'], i['tds'], 'Paid' if i['paid'] else 'Unpaid', i['received'] if i['paid'] else 0,
                i['received_date'].isoformat() if (i['paid'] and i['received_date']) else '', i['due'].isoformat() if i['due'] else '']
               for i in inv if _in(i['date'], a, b)], (6, 7, 8, 9, 11))
        sheet(wb.create_sheet('Expenses'), ['Date', 'Category', 'Type', 'Paid to', 'Amount (incl GST)', 'GST input', 'Cost', 'Status',
                                            'Mode', 'Client', 'Note'],
              [[x['d'].isoformat() if x['d'] else '', x['category'], x['kind'], x.get('payee') or '', x['amount'], x['gst_amount'],
                x['net_cost'], x['status'], x.get('payment_mode') or '', x.get('client_name') or '', x.get('note') or '']
               for x in ex if _in(x['d'], a, b)], (5, 6, 7))
        sheet(wb.create_sheet('GST'), ['Month', 'Output GST', 'Input GST', 'Estimated payable', 'Credit carried'],
              [[g['month'], g['output_gst'], g['input_gst'], g['payable'], g['credit_carry']] for g in gst_months(inv, ex, a, b)], (2, 3, 4, 5))
        sheet(wb.create_sheet('TDS'), ['FY', 'Client', 'TDS %', 'Invoiced', 'TDS deducted (paid)', 'TDS expected (unpaid)', 'Invoices'],
              [[t['fy'], t['client'], t['tds_percent'], t['invoiced'], t['deducted'], t['expected_unpaid'], t['invoices']]
               for t in tds_register(inv)], (4, 5, 6))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = 'HireLab_finance_%s_to_%s.xlsx' % (a.isoformat(), b.isoformat())
        return send_file(buf, as_attachment=True, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    finally:
        conn.close()


@pages.route('/finance', methods=['GET'])
def finance_page():
    core = _core()
    if not core.current_user():
        return redirect('/')
    p = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'finance.html')
    with open(p, encoding='utf-8') as f:
        return Response(f.read(), mimetype='text/html')

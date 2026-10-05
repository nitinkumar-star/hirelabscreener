"""
RecruitOS — Downloadable Call Analysis report  (Oct 2026)

One report builder, two outputs, two versions:

  GET /api/candidates/<cid>/call-analysis/<aid>/report
        ?format=docx | html        (html = print page -> "Save as PDF")
        &version=internal | client
        &transcript=1              (internal only)

  internal : everything on the Call Analysis tab (CTC, concerns, red flags,
             CV-vs-call checks, pitch follow-up, quotes, next step, transcript).
  client   : safe to send to the client — candidate snapshot, JD match,
             summary, requirement table, strengths, notice, interest.
             No CTC, red flags, concerns, internal checks, quotes, next step or
             transcript; any sentence that talks about salary is dropped.

Both versions end with the candidate's CV (from the profile): a PDF CV is
attached page by page as images, a Word / RTF / old .doc CV as its text.

The PDF is made by the browser (print page -> Save as PDF), so Hindi /
Devanagari transcripts print correctly without bundling fonts on the server.

ROLLBACK: remove 'call_report' from modules/__init__.py (the download buttons
then show an error; nothing else is affected).
"""

import io
import re
import json
import html as _h

from flask import Blueprint, jsonify, request, Response, send_file

from modules.shared import get_db, effective_company_id, login_required, current_user, _core

bp = Blueprint('call_report', __name__)

_MONEY = re.compile(r'(ctc|lpa|lakh|lac\b|lacs|salary|package|compensation|₹|\brs\.?\s*\d|inr|hike|increment|'
                    r'take[- ]home|in[- ]hand|per annum|\bcr\b|crore)', re.I)


def _client_text(s):
    """Drop every sentence that mentions pay. Returns '' if nothing is left."""
    s = str(s or '').strip()
    if not s:
        return ''
    parts = re.split(r'(?<=[.!?।])\s+', s)
    keep = [p for p in parts if not _MONEY.search(p)]
    return ' '.join(keep).strip()


def _list(v):
    return [x for x in (v if isinstance(v, list) else []) if x not in (None, '')]


def _status(v):
    return str(v or '').replace('_', ' ').strip().upper()


def _fmt_date(s):
    s = str(s or '')[:16].replace('T', ' ')
    try:
        import datetime
        d = datetime.datetime.strptime(s[:16], '%Y-%m-%d %H:%M')
        return d.strftime('%d %b %Y, %I:%M %p')
    except Exception:
        return s


def _agency_name():
    try:
        n = _core().get_setting('company_name', '') or ''
    except Exception:
        n = ''
    if not n:
        try:
            conn = get_db()
            r = conn.execute('SELECT name FROM companies WHERE id=?', (effective_company_id(),)).fetchone()
            conn.close()
            n = (r['name'] if r else '') or ''
        except Exception:
            n = ''
    return n


# ── CV attachment ─────────────────────────────────────────────────────────
CV_MAX_PAGES = 10


def load_cv(cand):
    """The candidate's CV, ready to append to the report, or None.
    {'name', 'pages': [jpeg bytes]}            for PDF CVs
    {'name', 'html': str, 'blocks': [...]}     for Word / RTF / .doc / Word-HTML CVs
    blocks: ('p', text, bold, style) | ('table', [[cell, ...], ...])"""
    import os
    core = _core()
    rel = os.path.basename(str(cand.get('cv_path') or '').strip())
    if not rel:
        return None
    fp = os.path.join(core.CV_DIR, rel)
    if not os.path.isfile(fp):
        return None
    name = str(cand.get('cv_original_name') or '').strip() or rel
    try:
        raw = open(fp, 'rb').read()
        kind = core._sniff_doc(raw)
    except Exception as e:
        print('[call-report] CV read failed:', e)
        return None
    try:
        if kind == 'pdf' or raw[:5] == b'%PDF-':
            return {'name': name, 'pages': _pdf_pages(raw), 'more': _pdf_count(raw) > CV_MAX_PAGES}
        if kind == 'zip':
            return {'name': name, **_docx_cv(core, raw)}
        if kind == 'ole':
            txt = core._doc_to_text(raw)
        elif kind == 'rtf':
            txt = core._rtf_to_text(raw)
        elif kind == 'html':
            txt = re.sub(r'<[^>]+>', '\n', core._html_doc_clean(raw) or '')
            txt = _h.unescape(re.sub(r'\n\s*\n+', '\n', txt))
        else:
            txt = ''
        lines = [l.strip() for l in (txt or '').splitlines() if l.strip()]
        if not lines:
            return {'name': name, 'html': '', 'blocks': [], 'unreadable': True}
        return {'name': name, 'blocks': [('p', l, False, '') for l in lines],
                'html': ''.join('<p>%s</p>' % _h.escape(l) for l in lines)}
    except Exception as e:
        print('[call-report] CV render failed:', e)
        return {'name': name, 'html': '', 'blocks': [], 'unreadable': True}


def _pdf_count(raw):
    try:
        import pypdfium2 as pdfium
        return len(pdfium.PdfDocument(raw))
    except Exception:
        return 0


def _pdf_pages(raw):
    import pypdfium2 as pdfium
    out = []
    pdf = pdfium.PdfDocument(raw)
    for i in range(min(len(pdf), CV_MAX_PAGES)):
        img = pdf[i].render(scale=2.0).to_pil().convert('RGB')     # ~144 dpi
        if img.width > 1400:
            img = img.resize((1400, int(img.height * 1400 / img.width)))
        b = io.BytesIO()
        img.save(b, 'JPEG', quality=80, optimize=True)
        out.append(b.getvalue())
    return out


def _docx_cv(core, raw):
    """Word CV -> HTML (mammoth, for the PDF page) + text blocks (for Word)."""
    clean, _ = core._docx_sanitize(raw)
    html = ''
    try:
        import mammoth
        try:
            html = mammoth.convert_to_html(io.BytesIO(clean)).value
        except Exception:
            fixed, n = core._docx_repair_xml(clean)
            html = mammoth.convert_to_html(io.BytesIO(fixed)).value if n else ''
            clean = fixed if n else clean
    except Exception:
        html = ''
    blocks = []
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = Document(io.BytesIO(clean))
        for el in d.element.body.iterchildren():
            tag = el.tag.split('}')[-1]
            if tag == 'p':
                p = Paragraph(el, d)
                t = p.text.strip()
                if t:
                    bold = bool(p.runs) and all((r.bold or False) for r in p.runs if r.text.strip())
                    sty = (p.style.name if p.style is not None else '') or ''
                    blocks.append(('p', t, bold, sty))
            elif tag == 'tbl':
                rows = []
                for r in Table(el, d).rows:
                    cells, seen = [], set()
                    for c in r.cells:
                        if id(c._tc) in seen:
                            continue                      # merged cell repeats
                        seen.add(id(c._tc)); cells.append(c.text.strip())
                    if any(cells):
                        rows.append(cells)
                if rows:
                    blocks.append(('table', rows))
    except Exception as e:
        print('[call-report] docx CV text failed:', e)
    if not html and blocks:
        html = ''.join('<p>%s</p>' % _h.escape(b[1]) for b in blocks if b[0] == 'p')
    return {'html': html, 'blocks': blocks, 'unreadable': not (html or blocks)}


# ── builder ───────────────────────────────────────────────────────────────
def build(conn, cid, rec, version='internal', transcript=False, with_cv=True):
    """A plain dict describing the report; renderers turn it into docx/html."""
    client = (version == 'client')
    cand = dict(conn.execute('SELECT * FROM candidates WHERE id=?', (cid,)).fetchone())
    m = None
    mid = rec['mandate_id'] or cand.get('mandate_id')
    if mid:
        m = conn.execute('SELECT role, client, location FROM mandates WHERE id=?', (mid,)).fetchone()
    try:
        a = json.loads(rec['analysis'] or '{}')
    except Exception:
        a = {}
    if not isinstance(a, dict):
        a = {}
    T = _client_text if client else (lambda s: str(s or '').strip())

    rep = {'version': version, 'agency': _agency_name(),
           'title': 'Candidate Assessment' if client else 'Call Analysis Report',
           'candidate': cand.get('name') or 'Candidate',
           'job': ' — '.join(x for x in ((m['role'] if m else ''), (m['client'] if m else '')) if x),
           'date': _fmt_date(rec['created_at']), 'by': rec['created_by_name'] or '',
           'facts': [], 'score': None, 'verdict': '', 'tiles': [], 'sections': []}

    # candidate snapshot
    exp = cand.get('experience')
    notice = a.get('notice_discussed_days')
    if notice in (None, ''):
        notice = cand.get('notice_period') or None
    facts = [('Current role', cand.get('designation')), ('Current company', cand.get('company')),
             ('Experience', f'{exp} yrs' if exp else ''), ('Location', cand.get('location'))]
    if notice not in (None, ''):
        nn = f'{notice} days'
        if a.get('notice_negotiable') is True:
            nn += ' (negotiable)'
        elif a.get('notice_negotiable') is False and not client:
            nn += ' (firm)'
        facts.append(('Notice period', nn))
    if not client:
        if a.get('ctc_discussed'):
            facts.append(('Current CTC (from call)', f"₹{a['ctc_discussed']} LPA"))
        if a.get('ctc_expected_discussed'):
            facts.append(('Expected CTC (from call)', f"₹{a['ctc_expected_discussed']} LPA"))
    rep['facts'] = [(k, str(v)) for k, v in facts if v not in (None, '', 0)]

    try:
        sc = a.get('jd_match_score')
        rep['score'] = max(0, min(100, int(float(sc)))) if sc not in (None, '') else None
    except (TypeError, ValueError):
        rep['score'] = None
    rep['verdict'] = a.get('jd_verdict') or a.get('fit_vs_jd') or ''
    if a.get('interest_level'):
        rep['tiles'].append(('Interest', str(a['interest_level']), T(a.get('interest_reason'))))
    if not client and a.get('recommendation'):
        rep['tiles'].append(('Recommendation', str(a['recommendation']), T(a.get('recommendation_reason'))))

    S = rep['sections']
    summ = T(a.get('overall_summary') or a.get('call_summary'))
    if summ:
        S.append({'title': 'Summary', 'kind': 'para', 'text': summ})
    if not client and a.get('overall_summary') and a.get('call_summary'):
        S.append({'title': 'On the call', 'kind': 'para', 'text': T(a.get('call_summary'))})

    reqs = [r for r in _list(a.get('requirements')) if isinstance(r, dict)]
    if reqs:
        S.append({'title': 'Job requirements — CV vs call', 'kind': 'table',
                  'head': ['Requirement', 'CV / profile', 'On the call', 'Status'],
                  'rows': [[T(r.get('requirement')) or '—', T(r.get('cv_evidence')) or '—',
                            T(r.get('call_evidence')) or '—', _status(r.get('status'))] for r in reqs]})

    strengths = [T(x) for x in _list(a.get('candidate_strengths'))]
    strengths = [x for x in strengths if x]
    if strengths:
        S.append({'title': 'Strengths', 'kind': 'bullets', 'items': strengths})

    if not client:
        for key, title in (('key_concerns', 'Key concerns'), ('red_flags', 'Red flags')):
            items = [str(x) for x in _list(a.get(key))]
            if items:
                S.append({'title': title, 'kind': 'bullets', 'items': items})
        cons = [c for c in _list(a.get('consistency')) if isinstance(c, dict)]
        if cons:
            S.append({'title': 'CV vs call — does it match?', 'kind': 'table',
                      'head': ['Topic', 'CV says', 'Call says', 'Status'],
                      'rows': [[c.get('topic') or '—', c.get('cv_says') or '—', c.get('call_says') or '—',
                                _status(c.get('status'))] for c in cons]})
        pf = [p for p in _list(a.get('pitch_followup')) if isinstance(p, dict)]
        if pf:
            S.append({'title': 'Pitch evaluation points — checked on the call', 'kind': 'table',
                      'head': ['Point', 'Note', 'Outcome'],
                      'rows': [[p.get('point') or '—', p.get('note') or '', _status(p.get('outcome'))] for p in pf]})
        quotes = [str(x) for x in _list(a.get('key_quotes'))]
        if quotes:
            S.append({'title': 'Key quotes from the candidate', 'kind': 'quotes', 'items': quotes})
        if a.get('next_step'):
            S.append({'title': 'Next step', 'kind': 'para',
                      'text': str(a['next_step']) + (f" — {a['next_step_deadline']}" if a.get('next_step_deadline') else '')})
        if transcript and (rec['transcript'] or '').strip():
            S.append({'title': 'Full transcript' + (f" ({a['languages_detected']})" if a.get('languages_detected') else ''),
                      'kind': 'pre', 'text': rec['transcript']})
    rep['cv'] = load_cv(cand) if with_cv else None
    return rep


def filename(rep, ext):
    base = re.sub(r'[^\w ()._-]+', '', f"{rep['candidate']} - Call Report"
                                                 f"{' (Client)' if rep['version'] == 'client' else ''}").strip()
    return (base or 'Call Report') + '.' + ext


# ── Word (.docx) ──────────────────────────────────────────────────────────
_COL = {'navy': '1F2A44', 'muted': '6B675E', 'green': '0F6E56', 'amber': 'A86A00', 'red': 'A32D2D', 'line': 'D9D5CC'}


def _status_color(s):
    s = (s or '').upper()
    if re.match(r'^(MET|MATCH|CONFIRMED|HIGH|PROCEED|STRONG)', s):
        return _COL['green']
    if re.match(r'^(GAP|MISMATCH|NOT CONFIRMED|LOW|REJECT|WEAK)', s):
        return _COL['red']
    return _COL['amber']


def to_docx(rep):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    rf.set(qn('w:cs'), 'Nirmala UI')          # Hindi / Devanagari text
    rf.set(qn('w:eastAsia'), 'Calibri')

    def run(p, text, size=None, bold=False, color=None, italic=False):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return r

    def para(space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        return p

    def heading(text):
        p = para(3)
        p.paragraph_format.space_before = Pt(10)
        run(p, text.upper(), 9, True, _COL['green'])

    def shade(cell, hexcol):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol)
        tcPr.append(sh)

    if rep['agency']:
        run(para(0), rep['agency'], 9, True, _COL['muted'])
    run(para(2), rep['title'], 18, True, _COL['navy'])
    p = para(2)
    run(p, rep['candidate'], 13, True)
    if rep['job']:
        run(p, '   ·   ' + rep['job'], 11, color=_COL['muted'])
    run(para(8), 'Call on ' + rep['date'] + (f" · by {rep['by']}" if rep['by'] else ''), 9, color=_COL['muted'])

    # headline row
    tiles = []
    if rep['score'] is not None:
        tiles.append(('JD match', f"{rep['score']}/100", rep['verdict']))
    elif rep['verdict']:
        tiles.append(('JD match', rep['verdict'], ''))
    tiles += rep['tiles']
    if tiles:
        t = doc.add_table(rows=1, cols=len(tiles))
        t.autofit = True
        for i, (lbl, val, sub) in enumerate(tiles):
            c = t.rows[0].cells[i]
            shade(c, 'F3F2EE')
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            run(c.paragraphs[0], lbl.upper(), 8, True, _COL['muted'])
            pv = c.add_paragraph(); pv.paragraph_format.space_after = Pt(0)
            run(pv, val, 15, True, _status_color(val) if lbl != 'JD match' else _COL['navy'])
            if sub:
                ps = c.add_paragraph(); ps.paragraph_format.space_after = Pt(2)
                run(ps, sub, 8.5, color=_COL['muted'])

    if rep['facts']:
        heading('Candidate snapshot')
        t = doc.add_table(rows=0, cols=2)
        t.style = 'Table Grid'
        for k, v in rep['facts']:
            r = t.add_row().cells
            shade(r[0], 'F6F5F1')
            run(r[0].paragraphs[0], k, 9.5, True, _COL['muted'])
            run(r[1].paragraphs[0], v, 10)
            r[0].width, r[1].width = Cm(5), Cm(12)

    for s in rep['sections']:
        heading(s['title'])
        k = s['kind']
        if k == 'para':
            run(para(4), s['text'])
        elif k == 'pre':
            for line in s['text'].splitlines() or ['']:
                run(para(1), line, 9.5)
        elif k == 'bullets':
            for it in s['items']:
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(1)
                run(p, it)
        elif k == 'quotes':
            for it in s['items']:
                p = para(3)
                p.paragraph_format.left_indent = Cm(0.6)
                run(p, f'“{it}”', italic=True, color='444444')
        elif k == 'table':
            t = doc.add_table(rows=1, cols=len(s['head']))
            t.style = 'Table Grid'
            for i, hd in enumerate(s['head']):
                c = t.rows[0].cells[i]
                shade(c, 'EAF4EF')
                run(c.paragraphs[0], hd, 9, True, _COL['navy'])
            for row in s['rows']:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    last = (i == len(row) - 1)
                    run(cells[i].paragraphs[0], str(v), 9.5, bold=last or i == 0,
                        color=_status_color(v) if last else None)

    p = para(0)
    p.paragraph_format.space_before = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _foot_p = p
    run(p, ('Prepared by ' + rep['agency'] if rep['agency'] else 'Generated by HireLab Screener')
        + (' · AI-assisted assessment from the screening call' if rep['version'] == 'client' else ' · Internal — do not forward'),
        8, color=_COL['muted'])

    cv = rep.get('cv')
    if cv:
        from docx.enum.text import WD_BREAK
        _foot_p.add_run().add_break(WD_BREAK.PAGE)
        p = para(6)
        run(p, 'ATTACHED: CV', 9, True, _COL['green'])
        run(p, '   ' + cv['name'], 9, color=_COL['muted'])
        if cv.get('pages'):
            for i, jpg in enumerate(cv['pages']):
                if i:
                    para(0).add_run().add_break(WD_BREAK.PAGE)
                pp = para(0)
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pp.add_run().add_picture(io.BytesIO(jpg), width=Cm(17))
            if cv.get('more'):
                run(para(0), f'(First {CV_MAX_PAGES} pages shown — open the full CV from the candidate profile.)',
                    8.5, color=_COL['muted'])
        elif cv.get('blocks'):
            for b in cv['blocks']:
                if b[0] == 'p':
                    sty = (b[3] or '').lower()
                    if sty.startswith('heading') or sty == 'title':
                        q = para(3); q.paragraph_format.space_before = Pt(8)
                        run(q, b[1], 11.5, True, _COL['navy'])
                    elif 'list' in sty:
                        q = doc.add_paragraph(style='List Bullet'); q.paragraph_format.space_after = Pt(1)
                        run(q, b[1], 10, b[2])
                    else:
                        run(para(3), b[1], 10, b[2])
                else:
                    ncol = max(len(r) for r in b[1])
                    t = doc.add_table(rows=0, cols=ncol)
                    t.style = 'Table Grid'
                    for r in b[1]:
                        cells = t.add_row().cells
                        for i, v in enumerate(r):
                            run(cells[i].paragraphs[0], v, 9.5)
                    para(2)
        else:
            run(para(0), 'The CV file could not be read — open it from the candidate profile.', 10, color=_COL['muted'])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ── print page (browser -> Save as PDF) ───────────────────────────────────
def to_html(rep, autoprint=True):
    e = lambda s: _h.escape(str(s or ''))

    def chip(s):
        c = _status_color(s)
        return f'<span class="chip" style="color:#{c};border-color:#{c}55;background:#{c}12">{e(s or "—")}</span>'

    tiles = []
    if rep['score'] is not None:
        tiles.append(f'<div class="tile"><div class="lb">JD match</div><div class="big">{rep["score"]}<small>/100</small></div>'
                     f'<div class="sub">{e(rep["verdict"])}</div></div>')
    elif rep['verdict']:
        tiles.append(f'<div class="tile"><div class="lb">JD match</div><div class="val">{e(rep["verdict"])}</div></div>')
    for lbl, val, sub in rep['tiles']:
        tiles.append(f'<div class="tile"><div class="lb">{e(lbl)}</div><div class="val" style="color:#{_status_color(val)}">{e(val)}</div>'
                     + (f'<div class="sub">{e(sub)}</div>' if sub else '') + '</div>')

    body = []
    if rep['facts']:
        body.append('<h2>Candidate snapshot</h2><table class="kv">' + ''.join(
            f'<tr><th>{e(k)}</th><td>{e(v)}</td></tr>' for k, v in rep['facts']) + '</table>')
    for s in rep['sections']:
        k = s['kind']
        h = f'<h2>{e(s["title"])}</h2>'
        if k == 'para':
            h += f'<p>{e(s["text"])}</p>'
        elif k == 'pre':
            h += f'<div class="pre">{e(s["text"])}</div>'
        elif k == 'bullets':
            h += '<ul>' + ''.join(f'<li>{e(x)}</li>' for x in s['items']) + '</ul>'
        elif k == 'quotes':
            h += ''.join(f'<blockquote>“{e(x)}”</blockquote>' for x in s['items'])
        elif k == 'table':
            h += ('<table class="grid"><thead><tr>' + ''.join(f'<th>{e(x)}</th>' for x in s['head']) + '</tr></thead><tbody>'
                  + ''.join('<tr>' + ''.join(
                      (f'<td class="st">{chip(v)}</td>' if i == len(r) - 1 else f'<td{" class=b" if i == 0 else ""}>{e(v)}</td>')
                      for i, v in enumerate(r)) + '</tr>' for r in s['rows']) + '</tbody></table>')
        body.append(h)

    foot = (('Prepared by ' + rep['agency']) if rep['agency'] else 'Generated by HireLab Screener') + \
           (' · AI-assisted assessment from the screening call' if rep['version'] == 'client' else ' · Internal — do not forward')
    title = filename(rep, 'pdf')[:-4]
    return f'''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{e(title)}</title>
<style>
@page {{ size: A4; margin: 14mm 13mm; }}
*{{box-sizing:border-box}}
body{{margin:0;background:#EEECE6;color:#1F1E1B;font:13px/1.55 -apple-system,"Segoe UI",Roboto,"Noto Sans","Noto Sans Devanagari","Nirmala UI",Arial,sans-serif}}
.bar{{position:sticky;top:0;background:#1F2A44;color:#fff;padding:10px 16px;display:flex;gap:10px;align-items:center;justify-content:space-between;z-index:2}}
.bar button{{background:#1D9E75;color:#fff;border:0;border-radius:8px;padding:9px 16px;font-size:14px;font-weight:700;cursor:pointer}}
.bar span{{font-size:12.5px;opacity:.85}}
.page{{max-width:820px;margin:18px auto;background:#fff;padding:30px 34px;border-radius:10px;box-shadow:0 2px 10px rgba(0,0,0,.08)}}
.ag{{font-size:11px;font-weight:700;letter-spacing:.06em;text-transform:uppercase;color:#6B675E}}
h1{{font-size:22px;margin:2px 0 6px;color:#1F2A44}}
.who{{font-size:16px;font-weight:700}} .who span{{font-weight:400;color:#6B675E}}
.meta{{font-size:11.5px;color:#6B675E;margin:3px 0 16px}}
.tiles{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-bottom:6px}}
.tile{{background:#F6F5F1;border-radius:9px;padding:10px 12px}}
.lb{{font-size:9.5px;letter-spacing:.06em;text-transform:uppercase;color:#6B675E;font-weight:700}}
.big{{font-size:28px;font-weight:800;color:#1F2A44;line-height:1.1}} .big small{{font-size:12px;color:#6B675E}}
.val{{font-size:18px;font-weight:800;margin-top:2px}} .sub{{font-size:11px;color:#55524A;margin-top:3px}}
h2{{font-size:11px;letter-spacing:.07em;text-transform:uppercase;color:#0F6E56;margin:20px 0 7px;border-bottom:1px solid #E4E0D6;padding-bottom:4px;break-after:avoid}}
p{{margin:0 0 6px}} ul{{margin:0;padding-left:18px}} li{{margin-bottom:3px}}
table{{width:100%;border-collapse:collapse;font-size:12px}}
.kv th{{text-align:left;width:34%;color:#6B675E;font-weight:600;padding:5px 8px 5px 0;vertical-align:top}}
.kv td{{padding:5px 0}} .kv tr+tr th,.kv tr+tr td{{border-top:1px solid #F0EDE6}}
.grid th{{text-align:left;background:#EAF4EF;color:#1F2A44;font-size:11px;padding:6px 8px}}
.grid td{{padding:6px 8px;border-top:1px solid #ECE8DF;vertical-align:top}} .grid td.b{{font-weight:600}}
.grid td.st{{text-align:right;white-space:nowrap}} tr{{break-inside:avoid}}
.chip{{font-size:9.5px;font-weight:800;letter-spacing:.03em;border:1px solid;border-radius:20px;padding:2px 8px}}
blockquote{{margin:0 0 6px;padding:3px 12px;border-left:3px solid #D9D5CC;font-style:italic;color:#444}}
.pre{{white-space:pre-wrap;font-size:11.5px;line-height:1.65;color:#333}}
.foot{{margin-top:24px;text-align:center;font-size:10px;color:#8A867C}}
.cv{{break-before:page;page-break-before:always}}
.cvh{{font-size:11px;letter-spacing:.07em;text-transform:uppercase;color:#0F6E56;font-weight:700;margin-bottom:10px}} .cvh span{{color:#6B675E;text-transform:none;letter-spacing:0;font-weight:400}}
.cvimg{{display:block;width:100%;height:auto;border:1px solid #E4E0D6}}
.cvimg+.cvimg{{break-before:page;page-break-before:always;margin-top:14px}}
.cvdoc{{font-size:12.5px;line-height:1.55}} .cvdoc table{{border-collapse:collapse;width:100%;margin:8px 0}}
.cvdoc td,.cvdoc th{{border:1px solid #ddd;padding:4px 6px;font-size:11.5px}} .cvdoc img{{max-width:100%}}
.cvdoc h1,.cvdoc h2,.cvdoc h3{{color:#1F2A44;margin:12px 0 6px;font-size:15px;text-transform:none;letter-spacing:0;border:0}}
@media print{{ body{{background:#fff}} .bar{{display:none}} .page{{margin:0;padding:0;box-shadow:none;max-width:none;border-radius:0}}
  .cvimg{{border:0;max-height:265mm;width:auto;max-width:100%;margin:0 auto}}
  .tile,.grid th,.chip{{-webkit-print-color-adjust:exact;print-color-adjust:exact}} }}
@media (max-width:600px){{ .page{{margin:0;border-radius:0;padding:20px 16px}} }}
</style></head><body>
<div class="bar"><span>Print dialog mein <b>Save as PDF</b> chunein</span><button onclick="window.print()">&#11015; Save as PDF</button></div>
<div class="page">
{f'<div class="ag">{e(rep["agency"])}</div>' if rep['agency'] else ''}
<h1>{e(rep['title'])}</h1>
<div class="who">{e(rep['candidate'])}{f' <span>· {e(rep["job"])}</span>' if rep['job'] else ''}</div>
<div class="meta">Call on {e(rep['date'])}{f' · by {e(rep["by"])}' if rep['by'] else ''}</div>
{('<div class="tiles">' + ''.join(tiles) + '</div>') if tiles else ''}
{''.join(body)}
<div class="foot">{e(foot)}</div>
</div>
{_cv_html(rep.get('cv'))}
{'<script>window.addEventListener("load",function(){setTimeout(function(){try{window.print()}catch(e){}},500)});</script>' if autoprint else ''}
</body></html>'''


def _cv_html(cv):
    if not cv:
        return ''
    import base64
    head = f'<div class="cvh">Attached: CV <span>· {_h.escape(cv["name"])}</span></div>'
    if cv.get('pages'):
        inner = ''.join(f'<img class="cvimg" alt="CV page {i + 1}" src="data:image/jpeg;base64,{base64.b64encode(j).decode()}">'
                        for i, j in enumerate(cv['pages']))
        if cv.get('more'):
            inner += f'<p style="font-size:11px;color:#6B675E">(First {CV_MAX_PAGES} pages shown — open the full CV from the candidate profile.)</p>'
    elif cv.get('html'):
        inner = '<div class="cvdoc">' + _strip_scripts(cv['html']) + '</div>'
    else:
        inner = '<p style="color:#6B675E">The CV file could not be read — open it from the candidate profile.</p>'
    return f'<div class="page cv">{head}{inner}</div>'


def _strip_scripts(h):
    h = re.sub(r'(?is)<(script|style|iframe|object|embed)[^>]*>.*?</\1>', '', h)
    h = re.sub(r'(?is)<(script|iframe|object|embed)[^>]*/?>', '', h)
    h = re.sub(r'(?i)\son\w+\s*=\s*("[^"]*"|\'[^\']*\'|[^\s>]+)', '', h)
    return re.sub(r'(?i)(href|src)\s*=\s*(["\']?)\s*javascript:', r'\1=\2#', h)


# ── endpoint ──────────────────────────────────────────────────────────────
@bp.route('/api/candidates/<int:cid>/call-analysis/<int:aid>/report', methods=['GET'])
@login_required
def call_report(cid, aid):
    fmt = (request.args.get('format') or 'html').lower()
    version = 'client' if (request.args.get('version') or '').lower() == 'client' else 'internal'
    want_tr = request.args.get('transcript') in ('1', 'true', 'yes') and version == 'internal'
    oid = effective_company_id()
    conn = get_db()
    own = conn.execute('SELECT owner_id FROM candidates WHERE id=?', (cid,)).fetchone()
    rec = conn.execute('SELECT * FROM candidate_call_analysis WHERE id=? AND candidate_id=? AND owner_id=?',
                       (aid, cid, oid)).fetchone() if own and own['owner_id'] == oid else None
    if not rec:
        conn.close(); return jsonify({'error': 'Not found'}), 404
    rep = build(conn, cid, rec, version, want_tr)
    conn.close()
    try:
        u = current_user() or {}
        _core().log_candidate_event(cid, 'call', f"Call report downloaded ({version}, {'Word' if fmt == 'docx' else 'PDF'})"
                                    + (f" by {u.get('display_name') or u.get('username')}" if u else ''))
    except Exception:
        pass
    if fmt == 'docx':
        return send_file(to_docx(rep), as_attachment=True, download_name=filename(rep, 'docx'),
                         mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    resp = Response(to_html(rep, autoprint=request.args.get('print', '1') != '0'), mimetype='text/html')
    resp.headers['Cache-Control'] = 'no-store'
    return resp
